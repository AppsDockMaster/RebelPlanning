#!/usr/bin/env python3
import io
import json
import os
import sqlite3
import subprocess
import tempfile
from datetime import datetime, timedelta
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlparse
from collections import defaultdict

try:
    import openpyxl
except Exception:
    openpyxl = None

BASE_DIR = Path(__file__).resolve().parent
# DATA_DIR: configurable via env var so Railway Volume can be mounted here
DATA_DIR = Path(os.environ.get("DATA_DIR", str(BASE_DIR)))
DATA_DIR.mkdir(parents=True, exist_ok=True)

DB_PATH = DATA_DIR / "planning.db"
TIMESHEET_PATHS = [
    DATA_DIR / "Employee Timesheets BE.xlsx",
    DATA_DIR / "Employee Timesheets NL.xlsx",
    DATA_DIR / "Employee Timesheet BE.xlsx",
    DATA_DIR / "Employee Timesheet NL.xlsx",
]
OHW_OPENING_PATHS = [
    DATA_DIR / "OHW 2025.xlsx",
    DATA_DIR / "OHW2025.xlsx",
]
INVOICE_LIST_PATHS = [
    DATA_DIR / "List of Invoices.xlsx",
    DATA_DIR / "List Of Invoices.xlsx",
]
PROJECT_LIST_PATHS = [
    DATA_DIR / "List of Projects.xlsx",
    DATA_DIR / "List Of Projects.xlsx",
]
VENDOR_INVOICE_PATHS = [
    DATA_DIR / "Vendor Invoices.xlsx",
    DATA_DIR / "Vendor invoices.xlsx",
]
HOST = "0.0.0.0"
PORT = int(os.environ.get("PORT", 8000))


def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    return conn


def init_db():
    with get_db() as conn:
        conn.executescript(
            """
            CREATE TABLE IF NOT EXISTS consultants (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL,
                start_date TEXT NOT NULL,
                entity TEXT NOT NULL DEFAULT 'RPL BE',
                exit_date TEXT
            );

            CREATE TABLE IF NOT EXISTS role_history (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                consultant_id INTEGER NOT NULL,
                role TEXT NOT NULL,
                from_date TEXT NOT NULL,
                to_date TEXT,
                FOREIGN KEY (consultant_id) REFERENCES consultants(id) ON DELETE CASCADE
            );

            CREATE TABLE IF NOT EXISTS regime_history (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                consultant_id INTEGER NOT NULL,
                regime INTEGER NOT NULL,
                from_date TEXT NOT NULL,
                to_date TEXT,
                FOREIGN KEY (consultant_id) REFERENCES consultants(id) ON DELETE CASCADE
            );

            CREATE TABLE IF NOT EXISTS consultant_cost_history (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                consultant_id INTEGER NOT NULL,
                year INTEGER NOT NULL,
                cost REAL NOT NULL,
                UNIQUE (consultant_id, year),
                FOREIGN KEY (consultant_id) REFERENCES consultants(id) ON DELETE CASCADE
            );

            CREATE TABLE IF NOT EXISTS consultant_vacations (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                consultant_id INTEGER NOT NULL,
                start_date TEXT NOT NULL,
                end_date TEXT NOT NULL,
                FOREIGN KEY (consultant_id) REFERENCES consultants(id) ON DELETE CASCADE
            );

            CREATE TABLE IF NOT EXISTS overhead_cost_history (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                year INTEGER NOT NULL,
                entity TEXT NOT NULL,
                cost REAL NOT NULL,
                UNIQUE (year, entity)
            );

            CREATE TABLE IF NOT EXISTS offers (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL,
                created_date TEXT NOT NULL DEFAULT (date('now')),
                submission_date TEXT NOT NULL,
                status TEXT NOT NULL DEFAULT 'open',
                close_reason TEXT
            );

            CREATE TABLE IF NOT EXISTS offer_staff (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                offer_id INTEGER NOT NULL,
                consultant_id INTEGER NOT NULL,
                UNIQUE (offer_id, consultant_id),
                FOREIGN KEY (offer_id) REFERENCES offers(id) ON DELETE CASCADE,
                FOREIGN KEY (consultant_id) REFERENCES consultants(id) ON DELETE CASCADE
            );

            CREATE TABLE IF NOT EXISTS offer_allocations (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                offer_id INTEGER NOT NULL,
                consultant_id INTEGER NOT NULL,
                week_start TEXT NOT NULL,
                mandays REAL NOT NULL,
                FOREIGN KEY (offer_id) REFERENCES offers(id) ON DELETE CASCADE,
                FOREIGN KEY (consultant_id) REFERENCES consultants(id) ON DELETE CASCADE
            );

            CREATE TABLE IF NOT EXISTS projects (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                project_number TEXT NOT NULL DEFAULT '',
                name TEXT NOT NULL,
                created_date TEXT NOT NULL DEFAULT (date('now')),
                start_date TEXT,
                delivery_date TEXT NOT NULL,
                budget REAL NOT NULL DEFAULT 0,
                status TEXT NOT NULL DEFAULT 'actief',
                lead_consultant_id INTEGER,
                billing_type TEXT NOT NULL DEFAULT 'regie',
                previous_budgets TEXT,
                next_budgets TEXT,
                previous_budget_total REAL NOT NULL DEFAULT 0,
                next_budget_total REAL NOT NULL DEFAULT 0,
                cap_amount REAL NOT NULL DEFAULT 0
            );

            CREATE TABLE IF NOT EXISTS project_staff (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                project_id INTEGER NOT NULL,
                consultant_id INTEGER NOT NULL,
                hourly_rate REAL NOT NULL DEFAULT 0,
                UNIQUE (project_id, consultant_id),
                FOREIGN KEY (project_id) REFERENCES projects(id) ON DELETE CASCADE,
                FOREIGN KEY (consultant_id) REFERENCES consultants(id) ON DELETE CASCADE
            );

            CREATE TABLE IF NOT EXISTS project_allocations (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                project_id INTEGER NOT NULL,
                consultant_id INTEGER NOT NULL,
                week_start TEXT NOT NULL,
                mandays REAL NOT NULL,
                FOREIGN KEY (project_id) REFERENCES projects(id) ON DELETE CASCADE,
                FOREIGN KEY (consultant_id) REFERENCES consultants(id) ON DELETE CASCADE
            );
            """
        )
        columns = [r["name"] for r in conn.execute("PRAGMA table_info(consultants)").fetchall()]
        if "entity" not in columns:
            conn.execute("ALTER TABLE consultants ADD COLUMN entity TEXT NOT NULL DEFAULT 'RPL BE'")
        if "exit_date" not in columns:
            conn.execute("ALTER TABLE consultants ADD COLUMN exit_date TEXT")
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS consultant_cost_history (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                consultant_id INTEGER NOT NULL,
                year INTEGER NOT NULL,
                cost REAL NOT NULL,
                UNIQUE (consultant_id, year),
                FOREIGN KEY (consultant_id) REFERENCES consultants(id) ON DELETE CASCADE
            )
            """
        )
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS overhead_cost_history (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                year INTEGER NOT NULL,
                entity TEXT NOT NULL,
                cost REAL NOT NULL,
                UNIQUE (year, entity)
            )
            """
        )
        offer_columns = [r["name"] for r in conn.execute("PRAGMA table_info(offers)").fetchall()]
        if "created_date" not in offer_columns:
            conn.execute("ALTER TABLE offers ADD COLUMN created_date TEXT NOT NULL DEFAULT '1970-01-01'")
            conn.execute("UPDATE offers SET created_date = COALESCE(created_date, date('now')) WHERE created_date = '1970-01-01'")
        project_columns = [r["name"] for r in conn.execute("PRAGMA table_info(projects)").fetchall()]
        if "created_date" not in project_columns:
            conn.execute("ALTER TABLE projects ADD COLUMN created_date TEXT NOT NULL DEFAULT '1970-01-01'")
            conn.execute("UPDATE projects SET created_date = COALESCE(created_date, date('now')) WHERE created_date = '1970-01-01'")
        if "budget" not in project_columns:
            conn.execute("ALTER TABLE projects ADD COLUMN budget REAL NOT NULL DEFAULT 0")
        if "status" not in project_columns:
            conn.execute("ALTER TABLE projects ADD COLUMN status TEXT NOT NULL DEFAULT 'actief'")
        if "project_number" not in project_columns:
            conn.execute("ALTER TABLE projects ADD COLUMN project_number TEXT NOT NULL DEFAULT ''")
        if "start_date" not in project_columns:
            conn.execute("ALTER TABLE projects ADD COLUMN start_date TEXT")
            conn.execute("UPDATE projects SET start_date = created_date WHERE start_date IS NULL OR start_date = ''")
        if "lead_consultant_id" not in project_columns:
            conn.execute("ALTER TABLE projects ADD COLUMN lead_consultant_id INTEGER")
        if "billing_type" not in project_columns:
            conn.execute("ALTER TABLE projects ADD COLUMN billing_type TEXT NOT NULL DEFAULT 'regie'")
        if "previous_budgets" not in project_columns:
            conn.execute("ALTER TABLE projects ADD COLUMN previous_budgets TEXT")
        if "next_budgets" not in project_columns:
            conn.execute("ALTER TABLE projects ADD COLUMN next_budgets TEXT")
        if "previous_budget_total" not in project_columns:
            conn.execute("ALTER TABLE projects ADD COLUMN previous_budget_total REAL NOT NULL DEFAULT 0")
        if "next_budget_total" not in project_columns:
            conn.execute("ALTER TABLE projects ADD COLUMN next_budget_total REAL NOT NULL DEFAULT 0")
        if "cap_amount" not in project_columns:
            conn.execute("ALTER TABLE projects ADD COLUMN cap_amount REAL NOT NULL DEFAULT 0")
        project_staff_columns = [r["name"] for r in conn.execute("PRAGMA table_info(project_staff)").fetchall()]
        if "hourly_rate" not in project_staff_columns:
            conn.execute("ALTER TABLE project_staff ADD COLUMN hourly_rate REAL NOT NULL DEFAULT 0")


def parse_iso_date(value):
    try:
        return datetime.strptime(value, "%Y-%m-%d")
    except Exception:
        return None


def parse_optional_budget_total(value):
    if value in (None, ""):
        return 0.0
    try:
        parsed = float(value)
    except Exception:
        raise ValueError("Budget moet numeriek zijn.")
    if parsed < 0:
        raise ValueError("Budget kan niet negatief zijn.")
    return parsed


def normalize_name(value):
    return " ".join((value or "").strip().lower().split())


def normalize_project_number(value):
    txt = str(value or "").strip()
    if not txt:
        return ""
    numeric_candidate = txt.replace(",", ".")
    try:
        num = float(numeric_candidate)
        if num.is_integer():
            return str(int(num))
    except Exception:
        pass
    return txt


def start_of_week_iso(date_obj):
    monday = date_obj - timedelta(days=date_obj.weekday())
    return monday.strftime("%Y-%m-%d")


def get_available_timesheet_paths():
    return [path for path in TIMESHEET_PATHS if path.exists()]


def get_first_existing_path(candidates):
    for path in candidates:
        if path.exists():
            return path
    return None


def get_planning_weeks():
    today = datetime.now().date()
    current_monday = today - timedelta(days=today.weekday())
    return [(current_monday + timedelta(weeks=i)).strftime("%Y-%m-%d") for i in range(13)]


def get_planning_weeks_for_consultant(conn, consultant_id):
    """Return all Mondays from the current week up to the last allocation week for this consultant."""
    today = datetime.now().date()
    current_monday = today - timedelta(days=today.weekday())
    start_iso = current_monday.strftime("%Y-%m-%d")

    rows = conn.execute(
        """
        SELECT week_start FROM project_allocations
        WHERE consultant_id = ?
          AND EXISTS (SELECT 1 FROM projects p WHERE p.id = project_id AND LOWER(p.status) = 'actief')
        UNION
        SELECT week_start FROM offer_allocations
        WHERE consultant_id = ?
          AND EXISTS (SELECT 1 FROM offers o WHERE o.id = offer_id AND LOWER(o.status) = 'open')
        """,
        (consultant_id, consultant_id),
    ).fetchall()

    all_week_isos = {r["week_start"] for r in rows}
    future_weeks = [w for w in all_week_isos if w >= start_iso]
    max_iso = max(future_weeks) if future_weeks else None

    if not max_iso:
        return [(current_monday + timedelta(weeks=i)).strftime("%Y-%m-%d") for i in range(13)]

    weeks = []
    current = current_monday
    max_date = datetime.strptime(max_iso, "%Y-%m-%d").date()
    while current <= max_date:
        weeks.append(current.strftime("%Y-%m-%d"))
        current += timedelta(weeks=1)
    return weeks


def planning_week_label(iso_date):
    dt = parse_iso_date(iso_date)
    if not dt:
        return iso_date
    week_num = dt.isocalendar()[1]
    month_abbr = ["jan", "feb", "mrt", "apr", "mei", "jun", "jul", "aug", "sep", "okt", "nov", "dec"][dt.month - 1]
    return f"W{week_num} ({dt.day:02d} {month_abbr})"


def count_vacation_mandays_in_week(week_start_iso, start_iso, end_iso):
    week_start = parse_iso_date(week_start_iso)
    start_date = parse_iso_date(start_iso)
    end_date = parse_iso_date(end_iso)
    if not all([week_start, start_date, end_date]):
        return 0
    week_end = week_start + timedelta(days=6)
    from_date = max(start_date, week_start)
    to_date = min(end_date, week_end)
    if from_date > to_date:
        return 0
    days = 0
    current = from_date
    while current <= to_date:
        if current.weekday() < 5:
            days += 1
        current += timedelta(days=1)
    return days


def get_regime_for_date(regime_history, iso_date_str):
    for entry in regime_history:
        from_d = entry["from"]
        to_d = entry.get("to")
        if from_d <= iso_date_str and (not to_d or to_d >= iso_date_str):
            return entry["regime"]
    return 100


def get_weekly_capacity_days_py(start_date, exit_date, regime_history, week_start_iso):
    if start_date > week_start_iso:
        return 0
    if exit_date and exit_date < week_start_iso:
        return 0
    regime = get_regime_for_date(regime_history, week_start_iso)
    return (regime / 100) * 5


def export_consultant_planning(conn, consultant_id, capacity_chart_b64, timeline_chart_b64):
    import io
    import base64

    if openpyxl is None:
        return None, "openpyxl is niet beschikbaar."

    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from openpyxl.drawing.image import Image as XLImage

    consultant = conn.execute(
        "SELECT id, name, start_date, exit_date, entity FROM consultants WHERE id = ?",
        (consultant_id,),
    ).fetchone()
    if consultant is None:
        return None, "Consultant niet gevonden."

    regime_history = [
        {"from": r["from_date"], "to": r["to_date"], "regime": r["regime"]}
        for r in conn.execute(
            "SELECT from_date, to_date, regime FROM regime_history WHERE consultant_id = ? ORDER BY from_date",
            (consultant_id,),
        ).fetchall()
    ]

    vacations = conn.execute(
        "SELECT id, start_date, end_date FROM consultant_vacations WHERE consultant_id = ? ORDER BY start_date",
        (consultant_id,),
    ).fetchall()

    weeks = get_planning_weeks_for_consultant(conn, consultant_id)

    # Aggregate project allocations
    project_allocs = conn.execute(
        """
        SELECT pa.week_start, pa.mandays, p.id AS project_id, p.name AS project_name,
               p.project_number
        FROM project_allocations pa
        JOIN projects p ON p.id = pa.project_id
        WHERE pa.consultant_id = ? AND LOWER(p.status) = 'actief'
        ORDER BY p.name, pa.week_start
        """,
        (consultant_id,),
    ).fetchall()

    # Aggregate offer allocations
    offer_allocs = conn.execute(
        """
        SELECT oa.week_start, oa.mandays, o.id AS offer_id, o.name AS offer_name
        FROM offer_allocations oa
        JOIN offers o ON o.id = oa.offer_id
        WHERE oa.consultant_id = ? AND LOWER(o.status) = 'open'
        ORDER BY o.name, oa.week_start
        """,
        (consultant_id,),
    ).fetchall()

    # Build lookup dicts
    projects_by_id = {}
    for row in project_allocs:
        pid = row["project_id"]
        if pid not in projects_by_id:
            num = row["project_number"] or ""
            projects_by_id[pid] = {
                "name": row["project_name"],
                "number": num,
                "display": f"{num} - {row['project_name']}" if num else row["project_name"],
                "weeks": {w: 0.0 for w in weeks},
            }
        if row["week_start"] in weeks:
            projects_by_id[pid]["weeks"][row["week_start"]] = float(row["mandays"] or 0)

    offers_by_id = {}
    for row in offer_allocs:
        oid = row["offer_id"]
        if oid not in offers_by_id:
            offers_by_id[oid] = {"name": row["offer_name"], "weeks": {w: 0.0 for w in weeks}}
        if row["week_start"] in weeks:
            offers_by_id[oid]["weeks"][row["week_start"]] = float(row["mandays"] or 0)

    vacation_rows = []
    for vac in vacations:
        vac_weeks = {}
        for w in weeks:
            d = count_vacation_mandays_in_week(w, vac["start_date"], vac["end_date"])
            if d > 0:
                vac_weeks[w] = d
        if vac_weeks:
            vacation_rows.append({"label": f"{vac['start_date']} – {vac['end_date']}", "weeks": vac_weeks})

    # ── Workbook ──────────────────────────────────────────────────────────────
    from openpyxl.formatting.rule import CellIsRule

    wb = openpyxl.Workbook()

    # ── Sheet 1 : Planning table ──────────────────────────────────────────────
    ws = wb.active
    ws.title = "Planning"

    hdr_font_white = Font(bold=True, color="FFFFFF")
    hdr_fill_dark = PatternFill("solid", fgColor="1F4E79")
    center = Alignment(horizontal="center")

    CAP_ROW = 3
    BEZ_ROW = 4
    FIRST_DATA_ROW = 5

    # Row 1 – consultant name
    ws["A1"] = consultant["name"]
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells(f"A1:{get_column_letter(2 + len(weeks))}1")

    # Row 2 – column headers
    ws["A2"] = "Type"
    ws["B2"] = "Naam"
    for col in (ws["A2"], ws["B2"]):
        col.font = hdr_font_white
        col.fill = hdr_fill_dark
    for i, w in enumerate(weeks):
        cell = ws.cell(row=2, column=3 + i, value=planning_week_label(w))
        cell.font = hdr_font_white
        cell.fill = hdr_fill_dark
        cell.alignment = center

    # Row 3 – Capacity (static values)
    ws.cell(row=CAP_ROW, column=1, value="Capaciteit").font = Font(bold=True)
    ws.cell(row=CAP_ROW, column=2, value="Max (mandagen)").font = Font(bold=True)
    for i, w in enumerate(weeks):
        cap = get_weekly_capacity_days_py(
            consultant["start_date"], consultant["exit_date"], regime_history, w
        )
        cell = ws.cell(row=CAP_ROW, column=3 + i, value=round(cap, 1))
        cell.alignment = center

    # Row 4 – Bezetting labels only (formulas written after data rows)
    ws.cell(row=BEZ_ROW, column=1, value="Totaal").font = Font(bold=True)
    ws.cell(row=BEZ_ROW, column=2, value="Bezetting").font = Font(bold=True)

    # Data rows (row 5+)
    data_row = FIRST_DATA_ROW

    # Projects section
    if projects_by_id:
        hdr_cell = ws.cell(row=data_row, column=1, value="Projecten")
        hdr_cell.font = Font(bold=True, color="1F4E79")
        hdr_cell.fill = PatternFill("solid", fgColor="DDEEFF")
        ws.merge_cells(f"A{data_row}:{get_column_letter(2 + len(weeks))}{data_row}")
        data_row += 1
        for proj in projects_by_id.values():
            ws.cell(row=data_row, column=1, value="Project")
            ws.cell(row=data_row, column=2, value=proj["display"])
            for i, w in enumerate(weeks):
                v = proj["weeks"].get(w, 0)
                cell = ws.cell(row=data_row, column=3 + i, value=round(v, 1) if v else None)
                cell.alignment = center
            data_row += 1

    # Offers section
    if offers_by_id:
        hdr_cell = ws.cell(row=data_row, column=1, value="Offertes")
        hdr_cell.font = Font(bold=True, color="7B3F00")
        hdr_cell.fill = PatternFill("solid", fgColor="FFF3E0")
        ws.merge_cells(f"A{data_row}:{get_column_letter(2 + len(weeks))}{data_row}")
        data_row += 1
        for offer in offers_by_id.values():
            ws.cell(row=data_row, column=1, value="Offerte")
            ws.cell(row=data_row, column=2, value=offer["name"])
            for i, w in enumerate(weeks):
                v = offer["weeks"].get(w, 0)
                cell = ws.cell(row=data_row, column=3 + i, value=round(v, 1) if v else None)
                cell.alignment = center
            data_row += 1

    # Vacations section (static – derived from vacation calendar)
    if vacation_rows:
        hdr_cell = ws.cell(row=data_row, column=1, value="Vakantie")
        hdr_cell.font = Font(bold=True, color="444444")
        hdr_cell.fill = PatternFill("solid", fgColor="EEEEEE")
        ws.merge_cells(f"A{data_row}:{get_column_letter(2 + len(weeks))}{data_row}")
        data_row += 1
        for vac in vacation_rows:
            ws.cell(row=data_row, column=1, value="Vakantie")
            ws.cell(row=data_row, column=2, value=vac["label"])
            for i, w in enumerate(weeks):
                v = vac["weeks"].get(w, 0)
                cell = ws.cell(row=data_row, column=3 + i, value=round(v, 1) if v else None)
                cell.alignment = center
            data_row += 1

    last_data_row = max(data_row - 1, FIRST_DATA_ROW)

    # Row 4 – Bezetting formulas (now that we know the full data range)
    for i in range(len(weeks)):
        col = 3 + i
        col_letter = get_column_letter(col)
        cap_ref = f"{col_letter}{CAP_ROW}"
        sum_range = f"{col_letter}{FIRST_DATA_ROW}:{col_letter}{last_data_row}"
        cell = ws.cell(row=BEZ_ROW, column=col)
        cell.value = f'=IF({cap_ref}>0,SUM({sum_range})/{cap_ref},"n.v.t.")'
        cell.alignment = center
        cell.number_format = "0%"
        cell.font = Font(bold=True)

    # Conditional formatting on bezetting row.
    # Differential formats (dxf) require bgColor, not fgColor.
    # openpyxl assigns priority 1 to the FIRST added rule (top of Rules Manager).
    # Add most-to-least restrictive so red (>=1) has the highest priority.
    bez_range = f"C{BEZ_ROW}:{get_column_letter(2 + len(weeks))}{BEZ_ROW}"
    ws.conditional_formatting.add(bez_range, CellIsRule(
        operator="greaterThanOrEqual", formula=["1"],
        fill=PatternFill(bgColor="FFE53935"),
        font=Font(color="FFFFFF"),
    ))
    ws.conditional_formatting.add(bez_range, CellIsRule(
        operator="greaterThanOrEqual", formula=["0.9"],
        fill=PatternFill(bgColor="FFFB8C00"),
        font=Font(color="FFFFFF"),
    ))
    ws.conditional_formatting.add(bez_range, CellIsRule(
        operator="greaterThanOrEqual", formula=["0.75"],
        fill=PatternFill(bgColor="FFFDD835"),
    ))
    ws.conditional_formatting.add(bez_range, CellIsRule(
        operator="greaterThan", formula=["0"],
        fill=PatternFill(bgColor="FF66BB6A"),
    ))

    # Column widths
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 42
    for i in range(len(weeks)):
        ws.column_dimensions[get_column_letter(3 + i)].width = 13

    # ── Sheet 2 : Capaciteit chart ────────────────────────────────────────────
    ws2 = wb.create_sheet("Capaciteit")
    ws2["A1"] = f"Capaciteit – {consultant['name']}"
    ws2["A1"].font = Font(bold=True, size=13)
    if capacity_chart_b64:
        try:
            raw = capacity_chart_b64.split(",", 1)[-1]
            img = XLImage(io.BytesIO(base64.b64decode(raw)))
            img.anchor = "A3"
            ws2.add_image(img)
        except Exception as exc:
            ws2["A3"] = f"Grafiek kon niet worden ingebed: {exc}"

    # ── Sheet 3 : Project tijdslijn ───────────────────────────────────────────
    ws3 = wb.create_sheet("Tijdslijn")
    ws3["A1"] = f"Project tijdslijn – {consultant['name']}"
    ws3["A1"].font = Font(bold=True, size=13)
    if timeline_chart_b64:
        try:
            raw = timeline_chart_b64.split(",", 1)[-1]
            img = XLImage(io.BytesIO(base64.b64decode(raw)))
            img.anchor = "A3"
            ws3.add_image(img)
        except Exception as exc:
            ws3["A3"] = f"Grafiek kon niet worden ingebed: {exc}"

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.read(), consultant["name"]


def week_fully_passed(week_start_iso):
    dt = parse_iso_date(week_start_iso)
    if dt is None:
        return False
    week_end = dt + timedelta(days=6)
    return week_end.date() < datetime.now().date()


def parse_timesheet_date(value):
    if isinstance(value, datetime):
        return value
    if isinstance(value, str):
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y"):
            try:
                return datetime.strptime(value, fmt)
            except Exception:
                continue
    return None


def parse_task_int(value):
    task_str = str(value or "").strip().lower().replace(",", ".")
    for token in task_str.split():
        try:
            return int(float(token))
        except Exception:
            continue
    try:
        return int(float(task_str))
    except Exception:
        return None


def classify_timesheet_hours(project_name, task_value):
    project_norm = " ".join(str(project_name or "").strip().lower().split())
    task_int = parse_task_int(task_value)

    if task_int == 100:
        return "billable"
    if task_int == 110:
        return "acquisitie"

    # Maak de herkenning robuuster voor varianten in projectnaam.
    is_internal_rebel = (
        ("internal" in project_norm and "rebel" in project_norm and "port" in project_norm)
        or "internal project rebel ports & logistics" in project_norm
    )
    if is_internal_rebel:
        if task_int in {401, 413}:
            return "acquisitie"
        return "administratie"

    # Fallback: niet expliciet geclassificeerd, maar meegenomen als administratie.
    return "administratie"


def is_absence_project(project_name):
    project_norm = " ".join(str(project_name or "").strip().lower().split())
    return project_norm.startswith("absence")


def get_display_project_name(project_name, category):
    """For internal Rebel projects, split display name by category."""
    project_norm = " ".join(str(project_name or "").strip().lower().split())
    is_internal_rebel = (
        ("internal" in project_norm and "rebel" in project_norm and "port" in project_norm)
        or "internal project rebel ports & logistics" in project_norm
    )
    if is_internal_rebel:
        label = "Acquisitie" if category == "acquisitie" else "Administratie"
        return f"Internal Project - {label}"
    return project_name


def build_timesheet_analysis(conn):
    if openpyxl is None:
        return {"months": [], "consultants": [], "rows": [], "monthlyTotals": [], "error": "openpyxl niet beschikbaar."}
    timesheet_paths = get_available_timesheet_paths()
    if not timesheet_paths:
        return {"months": [], "consultants": [], "rows": [], "monthlyTotals": [], "error": "Employee Timesheet BE/NL.xlsx niet gevonden."}

    consultant_names = [r["name"] for r in conn.execute("SELECT name FROM consultants").fetchall()]
    if not consultant_names:
        return {"months": [], "consultants": [], "rows": [], "monthlyTotals": []}
    consultant_lookup = {normalize_name(name): name for name in consultant_names}
    project_valuation_factor = {}
    for p in list_projects(conn):
        project_key = normalize_project_number(p.get("projectNumber"))
        if project_key:
            factor = float(p.get("timesheetValuationFactor") or 1.0)
            project_valuation_factor[project_key] = factor if factor > 0 else 1.0

    sheet_rows = []
    for path in timesheet_paths:
        wb = openpyxl.load_workbook(path, data_only=True)
        if "Exported Data" not in wb.sheetnames:
            continue
        ws = wb["Exported Data"]

        header_cells = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if not header_cells:
            continue
        headers = [str(h).strip() if h is not None else "" for h in header_cells]
        idx = {name: i for i, name in enumerate(headers)}
        idx_normalized = {"".join(ch for ch in name.lower() if ch.isalnum()): i for i, name in enumerate(headers)}
        needed = ["Employee Name", "Entry Date", "Ent. Time Qty", "SP", "Project Name", "Task"]
        if any(col not in idx for col in needed):
            continue
        project_no_idx = idx.get("Project No")
        if project_no_idx is None:
            for key in ("projectno", "projectnumber"):
                if key in idx_normalized:
                    project_no_idx = idx_normalized[key]
                    break
        sheet_rows.append((ws, idx, project_no_idx))

    if not sheet_rows:
        return {"months": [], "consultants": [], "rows": [], "monthlyTotals": [], "error": "Sheet 'Exported Data' of verplichte kolommen ontbreken in Employee Timesheet BE/NL.xlsx."}

    agg = defaultdict(lambda: {"hours": 0.0, "revenue": 0.0, "overrun": 0.0, "gross_revenue": 0.0})
    agg_week = defaultdict(lambda: {"hours": 0.0})
    agg_detail = defaultdict(lambda: {"hours": 0.0, "revenue": 0.0, "overrun": 0.0, "gross_revenue": 0.0})
    agg_project_totals = defaultdict(lambda: {"hours": 0.0, "revenue": 0.0, "overrun": 0.0, "gross_revenue": 0.0})
    agg_internal = defaultdict(lambda: {"hours": 0.0})
    month_totals = defaultdict(lambda: {"hours": 0.0, "revenue": 0.0, "overrun": 0.0, "gross_revenue": 0.0})
    mix_totals = defaultdict(float)
    mix_monthly = defaultdict(lambda: defaultdict(float))
    mix_consultant = defaultdict(lambda: defaultdict(float))
    mix_consultant_monthly = defaultdict(lambda: defaultdict(float))
    months_set = set()
    weeks_set = set()
    consultants_set = set()

    for ws, idx, project_no_idx in sheet_rows:
        for row in ws.iter_rows(min_row=2, values_only=True):
            raw_name = row[idx["Employee Name"]] if idx["Employee Name"] < len(row) else None
            raw_date = row[idx["Entry Date"]] if idx["Entry Date"] < len(row) else None
            raw_qty = row[idx["Ent. Time Qty"]] if idx["Ent. Time Qty"] < len(row) else None
            raw_sp = row[idx["SP"]] if idx["SP"] < len(row) else None
            raw_project = row[idx["Project Name"]] if idx["Project Name"] < len(row) else None
            raw_task = row[idx["Task"]] if idx["Task"] < len(row) else None
            raw_project_no = row[project_no_idx] if project_no_idx is not None and project_no_idx < len(row) else None

            canonical_name = consultant_lookup.get(normalize_name(str(raw_name or "")))
            if not canonical_name:
                continue
            date_obj = parse_timesheet_date(raw_date)
            if not date_obj:
                continue
            try:
                qty = float(raw_qty or 0)
                sp = float(raw_sp or 0)
            except Exception:
                continue
            if qty == 0:
                continue

            month = date_obj.strftime("%Y-%m")
            week_start = start_of_week_iso(date_obj)
            project_name = str(raw_project or "Onbekend project").strip() or "Onbekend project"
            project_no = normalize_project_number(raw_project_no)
            valuation_factor = project_valuation_factor.get(project_no, 1.0)
            category = classify_timesheet_hours(project_name, raw_task)
            exclude_from_mix = is_absence_project(project_name)
            key = (month, canonical_name)
            gross_revenue = qty * sp
            effective_revenue = gross_revenue * valuation_factor
            overrun_revenue = max(0.0, gross_revenue - effective_revenue)
            agg[key]["hours"] += qty
            agg[key]["revenue"] += effective_revenue
            agg[key]["overrun"] += overrun_revenue
            agg[key]["gross_revenue"] += gross_revenue
            week_key = (week_start, canonical_name)
            agg_week[week_key]["hours"] += qty
            display_name = get_display_project_name(project_name, category)
            detail_key = (month, canonical_name, display_name)
            agg_detail[detail_key]["hours"] += qty
            agg_detail[detail_key]["revenue"] += effective_revenue
            agg_detail[detail_key]["overrun"] += overrun_revenue
            agg_detail[detail_key]["gross_revenue"] += gross_revenue
            project_key = (canonical_name, display_name)
            agg_project_totals[project_key]["hours"] += qty
            agg_project_totals[project_key]["revenue"] += effective_revenue
            agg_project_totals[project_key]["overrun"] += overrun_revenue
            agg_project_totals[project_key]["gross_revenue"] += gross_revenue
            month_totals[month]["hours"] += qty
            month_totals[month]["revenue"] += effective_revenue
            month_totals[month]["overrun"] += overrun_revenue
            month_totals[month]["gross_revenue"] += gross_revenue
            if not exclude_from_mix:
                mix_totals[category] += qty
                mix_monthly[month][category] += qty
                mix_consultant[canonical_name][category] += qty
                mix_consultant_monthly[(canonical_name, month)][category] += qty
                if category != "billable":
                    task_int = parse_task_int(raw_task)
                    agg_internal[(month, canonical_name, task_int, category)]["hours"] += qty
            months_set.add(month)
            weeks_set.add(week_start)
            consultants_set.add(canonical_name)

    months = sorted(months_set)
    weeks = sorted(weeks_set)
    consultants = sorted(consultants_set)
    rows = []
    for month in months:
        for consultant in consultants:
            values = agg.get(
                (month, consultant),
                {"hours": 0.0, "revenue": 0.0, "overrun": 0.0, "gross_revenue": 0.0},
            )
            rows.append(
                {
                    "month": month,
                    "consultant": consultant,
                    "hours": round(values["hours"], 2),
                    "revenue": round(values["revenue"], 2),
                    "overrun": round(values["overrun"], 2),
                    "grossRevenue": round(values["gross_revenue"], 2),
                }
            )
    weekly_rows = []
    for week_start in weeks:
        for consultant in consultants:
            values = agg_week.get((week_start, consultant), {"hours": 0.0})
            weekly_rows.append(
                {
                    "weekStart": week_start,
                    "consultant": consultant,
                    "hours": round(values["hours"], 2),
                }
            )
    monthly_totals = [
        {
            "month": m,
            "hours": round(month_totals[m]["hours"], 2),
            "revenue": round(month_totals[m]["revenue"], 2),
            "overrun": round(month_totals[m]["overrun"], 2),
            "grossRevenue": round(month_totals[m]["gross_revenue"], 2),
        }
        for m in months
    ]
    detail_rows = [
        {
            "month": month,
            "consultant": consultant,
            "project": project,
            "hours": round(values["hours"], 2),
            "revenue": round(values["revenue"], 2),
            "overrun": round(values["overrun"], 2),
            "grossRevenue": round(values["gross_revenue"], 2),
        }
        for (month, consultant, project), values in sorted(agg_detail.items(), key=lambda x: (x[0][1], x[0][0], x[0][2]))
    ]
    project_totals = [
        {
            "consultant": consultant,
            "project": project,
            "hours": round(values["hours"], 2),
            "revenue": round(values["revenue"], 2),
            "overrun": round(values["overrun"], 2),
            "grossRevenue": round(values["gross_revenue"], 2),
        }
        for (consultant, project), values in sorted(agg_project_totals.items(), key=lambda x: (x[0][0], -x[1]["hours"], x[0][1]))
    ]
    mix_total_payload = {
        "billable": round(mix_totals["billable"], 2),
        "acquisitie": round(mix_totals["acquisitie"], 2),
        "administratie": round(mix_totals["administratie"], 2),
    }
    mix_monthly_rows = [
        {
            "month": m,
            "billable": round(mix_monthly[m]["billable"], 2),
            "acquisitie": round(mix_monthly[m]["acquisitie"], 2),
            "administratie": round(mix_monthly[m]["administratie"], 2),
        }
        for m in months
    ]
    mix_consultant_rows = [
        {
            "consultant": c,
            "billable": round(mix_consultant[c]["billable"], 2),
            "acquisitie": round(mix_consultant[c]["acquisitie"], 2),
            "administratie": round(mix_consultant[c]["administratie"], 2),
        }
        for c in consultants
    ]
    mix_consultant_monthly_rows = [
        {
            "consultant": c,
            "month": m,
            "billable": round(mix_consultant_monthly[(c, m)]["billable"], 2),
            "acquisitie": round(mix_consultant_monthly[(c, m)]["acquisitie"], 2),
            "administratie": round(mix_consultant_monthly[(c, m)]["administratie"], 2),
        }
        for c in consultants
        for m in months
    ]
    internal_task_labels = {
        400: "400 - Administration",
        401: "401 - Business development",
        402: "402 - Travel time (non chargeable)",
        403: "403 - General",
        404: "404 - Management",
        405: "405 - Team meeting",
        406: "406 - Knowledge sharing / -management and -development",
        407: "407 - Teambuilding",
        408: "408 - Management holding company",
        409: "409 - Housing",
        410: "410 - ISO9001",
        411: "411 - Recruitment and selection",
        412: "412 - Zeroca",
        413: "413 - Acquisition general",
        110: "110 - Acquisitie (task 110)",
    }
    internal_rows = [
        {
            "month": month,
            "consultant": consultant,
            "taskCode": task_code,
            "taskLabel": internal_task_labels.get(task_code, f"{task_code} - Overig" if task_code else "Onbekend"),
            "category": category,
            "hours": round(values["hours"], 2),
        }
        for (month, consultant, task_code, category), values in sorted(
            agg_internal.items(),
            key=lambda x: (x[0][2] or 0, x[0][3], x[0][0], x[0][1]),
        )
    ]
    return {
        "months": months,
        "weeks": weeks,
        "consultants": consultants,
        "rows": rows,
        "weeklyRows": weekly_rows,
        "monthlyTotals": monthly_totals,
        "detailRows": detail_rows,
        "projectTotals": project_totals,
        "internalRows": internal_rows,
        "mixTotals": mix_total_payload,
        "mixMonthly": mix_monthly_rows,
        "mixConsultant": mix_consultant_rows,
        "mixConsultantMonthly": mix_consultant_monthly_rows,
    }


def build_ohw_overview(conn):
    if openpyxl is None:
        return {"projects": [], "totals": {}, "error": "openpyxl niet beschikbaar."}

    ohw_opening_path = get_first_existing_path(OHW_OPENING_PATHS)
    invoice_path = get_first_existing_path(INVOICE_LIST_PATHS)
    timesheet_paths = get_available_timesheet_paths()
    if ohw_opening_path is None:
        return {"projects": [], "totals": {}, "error": "OHW 2025.xlsx niet gevonden."}
    if invoice_path is None:
        return {"projects": [], "totals": {}, "error": "List of Invoices.xlsx niet gevonden."}
    if not timesheet_paths:
        return {"projects": [], "totals": {}, "error": "Employee Timesheet BE/NL.xlsx niet gevonden."}
    project_list_path = get_first_existing_path(PROJECT_LIST_PATHS)

    projects = list_projects(conn)
    project_lookup = {
        normalize_project_number(p.get("projectNumber")): p for p in projects if normalize_project_number(p.get("projectNumber"))
    }
    external_project_names = {}
    external_project_managers = {}
    if project_list_path is not None:
        try:
            wb_projects = openpyxl.load_workbook(project_list_path, data_only=True)
            ws_projects = wb_projects["Exported Data"] if "Exported Data" in wb_projects.sheetnames else wb_projects[wb_projects.sheetnames[0]]
            header_cells = next(ws_projects.iter_rows(min_row=1, max_row=1, values_only=True), None)
            headers = [str(h).strip() if h is not None else "" for h in (header_cells or [])]
            idx = {name: i for i, name in enumerate(headers)}
            idx_norm = {"".join(ch for ch in name.lower() if ch.isalnum()): i for i, name in enumerate(headers)}
            project_idx = idx.get("Project No.")
            if project_idx is None:
                project_idx = idx_norm.get("projectno") or idx_norm.get("projectnumber")
            name_idx = idx.get("Project Name")
            if name_idx is None:
                name_idx = idx_norm.get("projectname")
            manager_idx = idx.get("Project Manager")
            if manager_idx is None:
                manager_idx = idx_norm.get("projectmanager")
            if project_idx is not None:
                for row in ws_projects.iter_rows(min_row=2, values_only=True):
                    project_no = normalize_project_number(row[project_idx] if project_idx < len(row) else None)
                    if not project_no:
                        continue
                    if name_idx is not None:
                        name = str(row[name_idx] if name_idx < len(row) and row[name_idx] is not None else "").strip()
                        if name and project_no not in external_project_names:
                            external_project_names[project_no] = name
                    if manager_idx is not None:
                        manager = str(row[manager_idx] if manager_idx < len(row) and row[manager_idx] is not None else "").strip()
                        if manager and project_no not in external_project_managers:
                            external_project_managers[project_no] = manager
        except Exception:
            external_project_names = {}
            external_project_managers = {}
    valuation_factor_by_project = {
        k: float(v.get("timesheetValuationFactor") or 1.0)
        for k, v in project_lookup.items()
    }

    consultant_rows_db = conn.execute("SELECT name, entity FROM consultants").fetchall()
    consultant_names = [r["name"] for r in consultant_rows_db]
    consultant_lookup = {normalize_name(name): name for name in consultant_names}
    consultant_entity_by_name = {str(r["name"]): (r["entity"] or "RPL BE") for r in consultant_rows_db}

    opening_by_project = defaultdict(float)
    wb_ohw = openpyxl.load_workbook(ohw_opening_path, data_only=True)
    ws_ohw = wb_ohw[wb_ohw.sheetnames[0]]
    header_cells = next(ws_ohw.iter_rows(min_row=1, max_row=1, values_only=True), None)
    headers = [str(h).strip() if h is not None else "" for h in (header_cells or [])]
    idx = {name: i for i, name in enumerate(headers)}
    idx_norm = {"".join(ch for ch in name.lower() if ch.isalnum()): i for i, name in enumerate(headers)}
    project_idx = idx.get("Projectnummer")
    if project_idx is None:
        project_idx = idx_norm.get("projectnummer") or idx_norm.get("projectno") or idx_norm.get("projectnumber")
    ohw_idx = idx.get("OHW")
    if ohw_idx is None:
        ohw_idx = idx_norm.get("ohw")
    if project_idx is None or ohw_idx is None:
        return {"projects": [], "totals": {}, "error": "Kolommen Projectnummer/OHW ontbreken in OHW 2025.xlsx."}
    for row in ws_ohw.iter_rows(min_row=2, values_only=True):
        project_no = normalize_project_number(row[project_idx] if project_idx < len(row) else None)
        if not project_no:
            continue
        try:
            opening_val = float(row[ohw_idx] if ohw_idx < len(row) else 0)
        except Exception:
            opening_val = 0.0
        opening_by_project[project_no] += opening_val

    invoiced_by_project = defaultdict(float)
    wb_inv = openpyxl.load_workbook(invoice_path, data_only=True)
    ws_inv = wb_inv["Exported Data"] if "Exported Data" in wb_inv.sheetnames else wb_inv[wb_inv.sheetnames[0]]
    header_cells = next(ws_inv.iter_rows(min_row=1, max_row=1, values_only=True), None)
    headers = [str(h).strip() if h is not None else "" for h in (header_cells or [])]
    idx = {name: i for i, name in enumerate(headers)}
    idx_norm = {"".join(ch for ch in name.lower() if ch.isalnum()): i for i, name in enumerate(headers)}
    project_idx = idx.get("Project No.")
    if project_idx is None:
        project_idx = idx_norm.get("projectno") or idx_norm.get("projectnumber")
    sales_idx = idx.get("Sales")
    if sales_idx is None:
        sales_idx = idx_norm.get("sales")
    total_idx = idx.get("Total")
    if total_idx is None:
        total_idx = idx_norm.get("total")
    dc_idx = idx.get("Debit/Credit")
    if dc_idx is None:
        dc_idx = idx_norm.get("debitcredit")
    if project_idx is None or (sales_idx is None and total_idx is None):
        return {"projects": [], "totals": {}, "error": "Kolommen Project No. en Sales/Total ontbreken in List of Invoices.xlsx."}
    for row in ws_inv.iter_rows(min_row=2, values_only=True):
        project_no = normalize_project_number(row[project_idx] if project_idx < len(row) else None)
        if not project_no:
            continue
        amount_raw = None
        if sales_idx is not None and sales_idx < len(row):
            amount_raw = row[sales_idx]
        if amount_raw in (None, "") and total_idx is not None and total_idx < len(row):
            amount_raw = row[total_idx]
        try:
            amount = float(amount_raw or 0)
        except Exception:
            amount = 0.0
        debit_credit = str(row[dc_idx] if dc_idx is not None and dc_idx < len(row) else "").strip().lower()
        sign = -1.0 if debit_credit == "credit" else 1.0
        invoiced_by_project[project_no] += sign * amount

    vendor_invoiced_by_project = defaultdict(float)
    vendor_invoice_path = get_first_existing_path(VENDOR_INVOICE_PATHS)
    if vendor_invoice_path is not None:
        wb_vend = openpyxl.load_workbook(vendor_invoice_path, data_only=True)
        ws_vend = wb_vend["Exported Data"] if "Exported Data" in wb_vend.sheetnames else wb_vend[wb_vend.sheetnames[0]]
        header_cells = next(ws_vend.iter_rows(min_row=1, max_row=1, values_only=True), None)
        headers = [str(h).strip() if h is not None else "" for h in (header_cells or [])]
        idx = {name: i for i, name in enumerate(headers)}
        idx_norm = {"".join(ch for ch in name.lower() if ch.isalnum()): i for i, name in enumerate(headers)}
        vend_project_idx = idx_norm.get("projectno")
        if vend_project_idx is None:
            vend_project_idx = idx_norm.get("projectnumber")
        amount_candidates = [
            "projectbalance", "projectbal", "balance", "total", "amount",
            "sales", "bedrag", "totaal",
        ]
        vend_total_idx = None
        for cand in amount_candidates:
            val = idx_norm.get(cand)
            if val is not None:
                vend_total_idx = val
                break
        if vend_total_idx is None:
            # Partial match: any header whose normalized name contains "balance" or starts with "projectbal"
            for norm_name, col_idx in idx_norm.items():
                if "balance" in norm_name or norm_name.startswith("projectbal"):
                    vend_total_idx = col_idx
                    break
        if vend_project_idx is not None and vend_total_idx is not None:
            for row in ws_vend.iter_rows(min_row=2, values_only=True):
                project_no = normalize_project_number(row[vend_project_idx] if vend_project_idx < len(row) else None)
                if not project_no:
                    continue
                try:
                    amount = float(row[vend_total_idx] if vend_total_idx < len(row) else 0)
                except Exception:
                    amount = 0.0
                vendor_invoiced_by_project[project_no] += -amount

    performed_by_project = defaultdict(float)
    performed_by_project_consultant = defaultdict(float)
    for path in timesheet_paths:
        wb = openpyxl.load_workbook(path, data_only=True)
        if "Exported Data" not in wb.sheetnames:
            continue
        ws = wb["Exported Data"]
        header_cells = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if not header_cells:
            continue
        headers = [str(h).strip() if h is not None else "" for h in header_cells]
        idx = {name: i for i, name in enumerate(headers)}
        idx_norm = {"".join(ch for ch in name.lower() if ch.isalnum()): i for i, name in enumerate(headers)}
        idx_employee = idx.get("Employee Name")
        if idx_employee is None:
            idx_employee = idx_norm.get("employeename")
        idx_qty = idx.get("Ent. Time Qty")
        if idx_qty is None:
            idx_qty = idx_norm.get("enttimeqty")
        idx_sp = idx.get("SP")
        if idx_sp is None:
            idx_sp = idx_norm.get("sp")
        idx_project_no = idx.get("Project No.")
        if idx_project_no is None:
            idx_project_no = idx_norm.get("projectno") or idx_norm.get("projectnumber")
        if None in {idx_employee, idx_qty, idx_sp, idx_project_no}:
            continue
        for row in ws.iter_rows(min_row=2, values_only=True):
            raw_name = row[idx_employee] if idx_employee < len(row) else None
            canonical_name = consultant_lookup.get(normalize_name(str(raw_name or "")))
            if not canonical_name:
                continue
            project_no = normalize_project_number(row[idx_project_no] if idx_project_no < len(row) else None)
            if not project_no:
                continue
            try:
                qty = float(row[idx_qty] if idx_qty < len(row) else 0)
                sp = float(row[idx_sp] if idx_sp < len(row) else 0)
            except Exception:
                continue
            if qty <= 0:
                continue
            factor = valuation_factor_by_project.get(project_no, 1.0)
            performed_amount = qty * sp * factor
            performed_by_project[project_no] += performed_amount
            performed_by_project_consultant[(project_no, canonical_name)] += performed_amount

    all_project_numbers = sorted(set(opening_by_project.keys()) | set(performed_by_project.keys()) | set(invoiced_by_project.keys()) | set(vendor_invoiced_by_project.keys()))
    rows = []
    totals = {"openingOHW": 0.0, "performedRevenue": 0.0, "invoiced": 0.0, "vendorInvoiced": 0.0, "currentOHW": 0.0}
    for project_no in all_project_numbers:
        opening = float(opening_by_project.get(project_no, 0.0))
        performed = float(performed_by_project.get(project_no, 0.0))
        invoiced = float(invoiced_by_project.get(project_no, 0.0))
        vendor_invoiced = float(vendor_invoiced_by_project.get(project_no, 0.0))
        current = opening + performed - invoiced + vendor_invoiced
        project = project_lookup.get(project_no, {})
        name = (project.get("name") if isinstance(project, dict) else "") or ""
        lead_name = (project.get("leadConsultantName") if isinstance(project, dict) else "") or ""
        if not name:
            name = external_project_names.get(project_no, "")
        if not lead_name:
            lead_name = external_project_managers.get(project_no, "")
        consultant_rows = [
            {
                "name": consultant_name,
                "entity": consultant_entity_by_name.get(consultant_name, ""),
                "performedRevenue": round(amount, 2),
            }
            for (proj_no, consultant_name), amount in performed_by_project_consultant.items()
            if proj_no == project_no and abs(float(amount)) > 0.000001
        ]
        consultant_rows.sort(key=lambda x: -float(x.get("performedRevenue") or 0))
        rows.append(
            {
                "projectNumber": project_no,
                "projectName": name,
                "projectLead": lead_name,
                "openingOHW": round(opening, 2),
                "performedRevenue": round(performed, 2),
                "invoiced": round(invoiced, 2),
                "vendorInvoiced": round(vendor_invoiced, 2),
                "currentOHW": round(current, 2),
                "consultants": consultant_rows,
            }
        )
        totals["openingOHW"] += opening
        totals["performedRevenue"] += performed
        totals["invoiced"] += invoiced
        totals["vendorInvoiced"] += vendor_invoiced
        totals["currentOHW"] += current

    rows.sort(key=lambda r: r["projectNumber"])
    return {
        "projects": rows,
        "totals": {k: round(v, 2) for k, v in totals.items()},
        "source": {
            "ohwFile": ohw_opening_path.name,
            "invoiceFile": invoice_path.name,
            "timesheets": [p.name for p in timesheet_paths],
        },
    }


def build_invoice_list():
    if openpyxl is None:
        return {"rows": [], "columns": [], "error": "openpyxl niet beschikbaar."}
    invoice_path = get_first_existing_path(INVOICE_LIST_PATHS)
    if invoice_path is None:
        return {"rows": [], "columns": [], "error": "List of Invoices.xlsx niet gevonden."}
    wb = openpyxl.load_workbook(invoice_path, data_only=True)
    ws = wb["Exported Data"] if "Exported Data" in wb.sheetnames else wb[wb.sheetnames[0]]
    header_cells = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if not header_cells:
        return {"rows": [], "columns": [], "error": "Geen headers gevonden in List of Invoices.xlsx."}
    EXCLUDED_COLUMNS = {"Debit/Credit", "Invoice Type", "Currency", "Tax", "Total", "Currency (unit)", "Tax, Currency", "Re. Invoice No."}
    all_columns = [str(h).strip() if h is not None else "" for h in header_cells]
    columns = [c for c in all_columns if c not in EXCLUDED_COLUMNS]
    col_indices = [i for i, c in enumerate(all_columns) if c not in EXCLUDED_COLUMNS]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if all(v is None for v in row):
            continue
        row_dict = {}
        for i, col in zip(col_indices, columns):
            val = row[i] if i < len(row) else None
            if hasattr(val, "isoformat"):
                val = val.isoformat()
            elif isinstance(val, (int, float)):
                pass
            elif val is not None:
                val = str(val)
            row_dict[col] = val
        rows.append(row_dict)
    return {"rows": rows, "columns": columns}


def ranges_overlap(a_from, a_to, b_from, b_to):
    af = parse_iso_date(a_from)
    bf = parse_iso_date(b_from)
    if af is None or bf is None:
        return False
    at = parse_iso_date(a_to) if a_to else datetime.max
    bt = parse_iso_date(b_to) if b_to else datetime.max
    return af <= bt and bf <= at


def validate_timeline(conn, table, consultant_id, from_date, to_date, ignore_entry_id=None):
    from_dt = parse_iso_date(from_date)
    if from_dt is None:
        return "Van-datum is verplicht en moet yyyy-mm-dd zijn."

    if to_date:
        to_dt = parse_iso_date(to_date)
        if to_dt is None:
            return "Tot-datum moet yyyy-mm-dd zijn."
        if to_dt < from_dt:
            return "Tot-datum kan niet voor Van-datum liggen."

    row = conn.execute("SELECT start_date FROM consultants WHERE id = ?", (consultant_id,)).fetchone()
    if row is None:
        return "Consultant niet gevonden."

    start_dt = parse_iso_date(row["start_date"])
    if start_dt and from_dt < start_dt:
        return "Van-datum kan niet voor indiensttreding liggen."

    if ignore_entry_id is None:
        existing = conn.execute(
            f"SELECT id, from_date, to_date FROM {table} WHERE consultant_id = ?",
            (consultant_id,),
        ).fetchall()
    else:
        existing = conn.execute(
            f"SELECT id, from_date, to_date FROM {table} WHERE consultant_id = ? AND id != ?",
            (consultant_id, ignore_entry_id),
        ).fetchall()

    for e in existing:
        if ranges_overlap(e["from_date"], e["to_date"], from_date, to_date):
            return "Per type mag een consultant geen overlappende periodes hebben."

    return None


def list_consultants(conn):
    consultants = conn.execute(
        "SELECT id, name, start_date, entity, exit_date FROM consultants ORDER BY name"
    ).fetchall()

    result = []
    for c in consultants:
        roles = conn.execute(
            """
            SELECT id, role, from_date, to_date
            FROM role_history
            WHERE consultant_id = ?
            ORDER BY from_date
            """,
            (c["id"],),
        ).fetchall()
        regimes = conn.execute(
            """
            SELECT id, regime, from_date, to_date
            FROM regime_history
            WHERE consultant_id = ?
            ORDER BY from_date
            """,
            (c["id"],),
        ).fetchall()
        costs = conn.execute(
            """
            SELECT id, year, cost
            FROM consultant_cost_history
            WHERE consultant_id = ?
            ORDER BY year
            """,
            (c["id"],),
        ).fetchall()
        vacations = conn.execute(
            """
            SELECT id, start_date, end_date
            FROM consultant_vacations
            WHERE consultant_id = ?
            ORDER BY start_date, end_date
            """,
            (c["id"],),
        ).fetchall()

        result.append(
            {
                "id": c["id"],
                "name": c["name"],
                "startDate": c["start_date"],
                "entity": c["entity"],
                "exitDate": c["exit_date"] or "",
                "roleHistory": [
                    {
                        "id": r["id"],
                        "role": r["role"],
                        "from": r["from_date"],
                        "to": r["to_date"] or "",
                    }
                    for r in roles
                ],
                "regimeHistory": [
                    {
                        "id": r["id"],
                        "regime": r["regime"],
                        "from": r["from_date"],
                        "to": r["to_date"] or "",
                    }
                    for r in regimes
                ],
                "costHistory": [
                    {
                        "id": row["id"],
                        "year": int(row["year"]),
                        "cost": float(row["cost"] or 0),
                    }
                    for row in costs
                ],
                "vacationHistory": [
                    {
                        "id": row["id"],
                        "startDate": row["start_date"],
                        "endDate": row["end_date"],
                    }
                    for row in vacations
                ],
            }
        )
    return result


def list_overhead_costs(conn):
    rows = conn.execute(
        """
        SELECT id, year, entity, cost
        FROM overhead_cost_history
        ORDER BY year, entity
        """
    ).fetchall()
    return [
        {
            "id": row["id"],
            "year": int(row["year"]),
            "entity": row["entity"],
            "cost": float(row["cost"] or 0),
        }
        for row in rows
    ]


def list_offers(conn):
    offers = conn.execute(
        "SELECT id, name, created_date, submission_date, status, close_reason FROM offers ORDER BY submission_date DESC, id DESC"
    ).fetchall()

    result = []
    for offer in offers:
        staff = conn.execute(
            """
            SELECT os.consultant_id, c.name AS consultant_name
            FROM offer_staff os
            JOIN consultants c ON c.id = os.consultant_id
            WHERE os.offer_id = ?
            ORDER BY c.name
            """,
            (offer["id"],),
        ).fetchall()
        allocs = conn.execute(
            """
            SELECT oa.id, oa.consultant_id, c.name AS consultant_name, oa.week_start, oa.mandays
            FROM offer_allocations oa
            JOIN consultants c ON c.id = oa.consultant_id
            WHERE oa.offer_id = ?
            ORDER BY oa.week_start, c.name
            """,
            (offer["id"],),
        ).fetchall()
        result.append(
            {
                "id": offer["id"],
                "name": offer["name"],
                "createdDate": offer["created_date"],
                "submissionDate": offer["submission_date"],
                "status": offer["status"],
                "closeReason": offer["close_reason"] or "",
                "staff": [
                    {"consultantId": s["consultant_id"], "consultantName": s["consultant_name"]}
                    for s in staff
                ],
                "allocations": [
                    {
                        "id": a["id"],
                        "consultantId": a["consultant_id"],
                        "consultantName": a["consultant_name"],
                        "weekStart": a["week_start"],
                        "mandays": float(a["mandays"]),
                    }
                    for a in allocs
                ],
            }
        )
    return result


def load_timesheet_project_actuals(conn):
    timesheet_paths = get_available_timesheet_paths()
    if openpyxl is None or not timesheet_paths:
        return {
            "week_mandays": {},
            "rate_by_project_consultant": {},
            "project_consultants": defaultdict(set),
            "consultant_names": {},
            "earliest_start_by_project": {},
        }
    sheet_rows = []
    for path in timesheet_paths:
        wb = openpyxl.load_workbook(path, data_only=True)
        if "Exported Data" not in wb.sheetnames:
            continue
        ws = wb["Exported Data"]
        header_cells = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if not header_cells:
            continue

        headers = [str(h).strip() if h is not None else "" for h in header_cells]

        def normalized_header(value):
            return "".join(ch for ch in str(value or "").lower() if ch.isalnum())

        idx_by_norm = {normalized_header(name): i for i, name in enumerate(headers)}

        def find_idx(*keys):
            for key in keys:
                if key in idx_by_norm:
                    return idx_by_norm[key]
            return None

        idx_employee = find_idx("employeename")
        idx_project_no = find_idx("projectno", "projectnumber")
        idx_entry_date = find_idx("entrydate")
        idx_qty = find_idx("enttimeqty")
        idx_sp = find_idx("sp")
        idx_task = find_idx("task")
        if None in {idx_employee, idx_project_no, idx_entry_date, idx_qty, idx_sp, idx_task}:
            continue
        sheet_rows.append((ws, idx_employee, idx_project_no, idx_entry_date, idx_qty, idx_sp, idx_task))

    if not sheet_rows:
        return {
            "week_mandays": {},
            "rate_by_project_consultant": {},
            "project_consultants": defaultdict(set),
            "consultant_names": {},
            "earliest_start_by_project": {},
        }

    consultant_rows = conn.execute("SELECT id, name FROM consultants").fetchall()
    consultant_lookup = {normalize_name(row["name"]): int(row["id"]) for row in consultant_rows}
    consultant_names = {int(row["id"]): row["name"] for row in consultant_rows}

    week_mandays = defaultdict(float)
    rate_agg = defaultdict(lambda: {"hours": 0.0, "revenue": 0.0})
    project_consultants = defaultdict(set)
    earliest_start_by_project = {}

    for ws, idx_employee, idx_project_no, idx_entry_date, idx_qty, idx_sp, idx_task in sheet_rows:
        for row in ws.iter_rows(min_row=2, values_only=True):
            raw_name = row[idx_employee] if idx_employee < len(row) else None
            raw_project_no = row[idx_project_no] if idx_project_no < len(row) else None
            raw_date = row[idx_entry_date] if idx_entry_date < len(row) else None
            raw_qty = row[idx_qty] if idx_qty < len(row) else None
            raw_sp = row[idx_sp] if idx_sp < len(row) else None
            raw_task = row[idx_task] if idx_task < len(row) else None

            consultant_id = consultant_lookup.get(normalize_name(str(raw_name or "")))
            if consultant_id is None:
                continue
            project_no = normalize_project_number(raw_project_no)
            if not project_no:
                continue
            date_obj = parse_timesheet_date(raw_date)
            if not date_obj:
                continue
            task_int = parse_task_int(raw_task)
            if task_int != 100:
                continue
            try:
                qty_hours = float(raw_qty or 0)
                sp = float(raw_sp or 0)
            except Exception:
                continue
            if qty_hours <= 0:
                continue

            week_start = start_of_week_iso(date_obj)
            entry_iso = date_obj.strftime("%Y-%m-%d")
            mandays = qty_hours / 8.0
            week_mandays[(project_no, consultant_id, week_start)] += mandays
            rate_agg[(project_no, consultant_id)]["hours"] += qty_hours
            rate_agg[(project_no, consultant_id)]["revenue"] += qty_hours * sp
            project_consultants[project_no].add(consultant_id)
            prev_start = earliest_start_by_project.get(project_no)
            if prev_start is None or entry_iso < prev_start:
                earliest_start_by_project[project_no] = entry_iso

    rate_by_project_consultant = {}
    for key, values in rate_agg.items():
        if values["hours"] > 0:
            rate_by_project_consultant[key] = values["revenue"] / values["hours"]

    return {
        "week_mandays": dict(week_mandays),
        "rate_by_project_consultant": rate_by_project_consultant,
        "project_consultants": project_consultants,
        "consultant_names": consultant_names,
        "earliest_start_by_project": earliest_start_by_project,
    }


def list_projects(conn):
    current_year = datetime.now().year
    timesheet_actuals = load_timesheet_project_actuals(conn)
    projects = conn.execute(
        """
        SELECT p.id, p.project_number, p.name, p.created_date, p.start_date, p.delivery_date, p.budget, p.status,
               p.lead_consultant_id, p.billing_type, p.previous_budget_total, p.next_budget_total, p.cap_amount,
               c.name AS lead_consultant_name
        FROM projects p
        LEFT JOIN consultants c ON c.id = p.lead_consultant_id
        ORDER BY p.delivery_date DESC, p.id DESC
        """
    ).fetchall()
    result = []
    for project in projects:
        project_number_norm = normalize_project_number(project["project_number"])
        raw_start_date = project["start_date"] or project["created_date"]
        timesheet_earliest_start = timesheet_actuals["earliest_start_by_project"].get(project_number_norm)
        effective_project_start = raw_start_date
        if timesheet_earliest_start and timesheet_earliest_start < effective_project_start:
            effective_project_start = timesheet_earliest_start
        db_staff_rows = conn.execute(
            """
            SELECT ps.consultant_id, c.name AS consultant_name, ps.hourly_rate
            FROM project_staff ps
            JOIN consultants c ON c.id = ps.consultant_id
            WHERE ps.project_id = ?
            ORDER BY c.name
            """,
            (project["id"],),
        ).fetchall()
        db_alloc_rows = conn.execute(
            """
            SELECT pa.id, pa.consultant_id, c.name AS consultant_name, pa.week_start, pa.mandays
            FROM project_allocations pa
            JOIN consultants c ON c.id = pa.consultant_id
            WHERE pa.project_id = ?
            ORDER BY pa.week_start, c.name
            """,
            (project["id"],),
        ).fetchall()
        staff_by_consultant = {}
        for s in db_staff_rows:
            consultant_id = int(s["consultant_id"])
            staff_by_consultant[consultant_id] = {
                "consultant_id": consultant_id,
                "consultant_name": s["consultant_name"],
                "hourly_rate": float(s["hourly_rate"] or 0),
            }
        for consultant_id in timesheet_actuals["project_consultants"].get(project_number_norm, set()):
            if consultant_id in staff_by_consultant:
                continue
            consultant_name = timesheet_actuals["consultant_names"].get(consultant_id)
            if not consultant_name:
                continue
            hourly_rate = float(timesheet_actuals["rate_by_project_consultant"].get((project_number_norm, consultant_id), 0.0))
            staff_by_consultant[consultant_id] = {
                "consultant_id": consultant_id,
                "consultant_name": consultant_name,
                "hourly_rate": hourly_rate,
            }
        staff = sorted(staff_by_consultant.values(), key=lambda x: (x["consultant_name"] or "").lower())
        rate_by_consultant = {int(s["consultant_id"]): float(s["hourly_rate"] or 0) for s in staff}

        alloc_by_key = {}
        for a in db_alloc_rows:
            consultant_id = int(a["consultant_id"])
            week_start = a["week_start"] or ""
            if week_fully_passed(week_start):
                continue
            alloc_by_key[(consultant_id, week_start)] = {
                "id": a["id"],
                "consultant_id": consultant_id,
                "consultant_name": a["consultant_name"],
                "week_start": week_start,
                "mandays": float(a["mandays"] or 0.0),
            }
        for (project_no, consultant_id, week_start), mandays in timesheet_actuals["week_mandays"].items():
            if project_no != project_number_norm:
                continue
            if not week_fully_passed(week_start):
                continue
            consultant_name = staff_by_consultant.get(consultant_id, {}).get("consultant_name") or timesheet_actuals["consultant_names"].get(consultant_id)
            if not consultant_name:
                continue
            alloc_by_key[(consultant_id, week_start)] = {
                "id": None,
                "consultant_id": consultant_id,
                "consultant_name": consultant_name,
                "week_start": week_start,
                "mandays": float(mandays or 0.0),
            }
        allocs = sorted(
            alloc_by_key.values(),
            key=lambda x: (x["week_start"] or "", (x["consultant_name"] or "").lower()),
        )
        expected_revenue_current_year = 0.0
        expected_revenue_previous_years = 0.0
        expected_revenue_next_years = 0.0
        expected_revenue_total = 0.0
        for a in allocs:
            week_start = a["week_start"] or ""
            week_year = None
            if len(str(week_start)) >= 4 and str(week_start)[:4].isdigit():
                week_year = int(str(week_start)[:4])
            consultant_id = int(a["consultant_id"])
            hourly_rate = rate_by_consultant.get(consultant_id, 0.0)
            mandays = float(a["mandays"] or 0.0)
            # allocaties worden als mandagen ingegeven; omzet = uren (mandagen * 8) * uurtarief.
            amount = (mandays * 8.0 * hourly_rate)
            expected_revenue_total += amount
            if week_year is None:
                continue
            if week_year < current_year:
                expected_revenue_previous_years += amount
            elif week_year > current_year:
                expected_revenue_next_years += amount
            else:
                expected_revenue_current_year += amount

        billing_type = project["billing_type"] or "regie"
        cap_amount = float(project["cap_amount"] or 0)
        if billing_type == "regie":
            total_budget = expected_revenue_total
            prev_budget_total = expected_revenue_previous_years
            next_budget_total = expected_revenue_next_years
        elif billing_type == "regie_cap":
            total_budget = cap_amount
            prev_budget_total = float(project["previous_budget_total"] or 0)
            next_budget_total = float(project["next_budget_total"] or 0)
        else:
            total_budget = float(project["budget"] or 0)
            prev_budget_total = float(project["previous_budget_total"] or 0)
            next_budget_total = float(project["next_budget_total"] or 0)
        current_year_budget = total_budget - prev_budget_total - next_budget_total
        expected_revenue_current_year_display = expected_revenue_current_year
        if billing_type in {"regie_cap", "fixed"}:
            expected_revenue_current_year_display = min(
                expected_revenue_current_year,
                max(current_year_budget, 0.0),
            )
        if current_year_budget > 0:
            if billing_type == "regie_cap":
                # Voor regie met cap meten we realisatie als prestaties t.o.v. factureerbare omzet
                # (omzet is gecapt op huidig jaarbudget).
                if expected_revenue_current_year_display > 0:
                    expected_realization_rate = expected_revenue_current_year / expected_revenue_current_year_display
                else:
                    expected_realization_rate = 0.0
            else:
                expected_realization_rate = expected_revenue_current_year / current_year_budget
        else:
            expected_realization_rate = 0.0
        timesheet_valuation_factor = 1.0
        if expected_realization_rate > 1:
            timesheet_valuation_factor = 1.0 / expected_realization_rate

        result.append(
            {
                "id": project["id"],
                "projectNumber": project["project_number"] or "",
                "name": project["name"],
                "createdDate": project["created_date"],
                "startDate": effective_project_start,
                "deliveryDate": project["delivery_date"],
                "budget": total_budget,
                "status": project["status"] or "actief",
                "leadConsultantId": project["lead_consultant_id"],
                "leadConsultantName": project["lead_consultant_name"] or "",
                "billingType": billing_type,
                "previousBudgetTotal": prev_budget_total,
                "nextBudgetTotal": next_budget_total,
                "capAmount": cap_amount,
                "currentYearBudget": round(current_year_budget, 2),
                "expectedRevenueCurrentYear": round(expected_revenue_current_year_display, 2),
                "expectedPerformanceCurrentYear": round(expected_revenue_current_year, 2),
                "expectedRealizationRate": round(expected_realization_rate, 4),
                "timesheetValuationFactor": round(timesheet_valuation_factor, 6),
                "staff": [
                    {
                        "consultantId": s["consultant_id"],
                        "consultantName": s["consultant_name"],
                        "hourlyRate": float(s["hourly_rate"] or 0),
                    }
                    for s in staff
                ],
                "allocations": [
                    {
                        "id": a["id"],
                        "consultantId": a["consultant_id"],
                        "consultantName": a["consultant_name"],
                        "weekStart": a["week_start"],
                        "mandays": float(a["mandays"]),
                    }
                    for a in allocs
                ],
            }
        )
    return result


class Handler(SimpleHTTPRequestHandler):
    def end_headers(self):
        self.send_header("Access-Control-Allow-Origin", "*")
        self.send_header("Access-Control-Allow-Methods", "GET, POST, OPTIONS")
        self.send_header("Access-Control-Allow-Headers", "Content-Type")
        super().end_headers()

    def json_response(self, status, payload):
        body = json.dumps(payload).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def read_json_body(self):
        length = int(self.headers.get("Content-Length", "0"))
        if length <= 0:
            return {}
        raw = self.rfile.read(length)
        return json.loads(raw.decode("utf-8"))

    def do_GET(self):
        parsed = urlparse(self.path)
        # Serve index.html for root, injecting a <base> tag when BASE_PATH is set
        if parsed.path in ("/", ""):
            base_path = os.environ.get("BASE_PATH", "")
            index_path = BASE_DIR / "index.html"
            try:
                html = index_path.read_text(encoding="utf-8")
                if base_path:
                    base_no_slash = base_path.rstrip("/")
                    html = html.replace(
                        "<head>",
                        f'<head>\n  <base href="{base_path}">\n  <script>window.PLANNING_BASE_PATH="{base_no_slash}";</script>',
                        1
                    )
                body = html.encode("utf-8")
                self.send_response(200)
                self.send_header("Content-Type", "text/html; charset=utf-8")
                self.send_header("Content-Length", str(len(body)))
                self.end_headers()
                self.wfile.write(body)
            except OSError:
                self.send_error(404)
            return
        if parsed.path == "/api/invoices":
            data = build_invoice_list()
            self.json_response(200, data)
            return
        if parsed.path == "/api/consultants":
            with get_db() as conn:
                data = list_consultants(conn)
                overhead = list_overhead_costs(conn)
            self.json_response(200, {"consultants": data, "overheadCosts": overhead})
            return
        if parsed.path == "/api/offers":
            with get_db() as conn:
                data = list_offers(conn)
            self.json_response(200, {"offers": data})
            return
        if parsed.path == "/api/projects":
            with get_db() as conn:
                data = list_projects(conn)
            self.json_response(200, {"projects": data})
            return
        if parsed.path == "/api/analysis/timesheets":
            with get_db() as conn:
                data = build_timesheet_analysis(conn)
            self.json_response(200, data)
            return
        if parsed.path == "/api/ohw":
            with get_db() as conn:
                data = build_ohw_overview(conn)
            self.json_response(200, data)
            return
        return super().do_GET()

    def do_OPTIONS(self):
        self.send_response(204)
        self.end_headers()

    def do_POST(self):
        parsed = urlparse(self.path)
        parts = [p for p in parsed.path.split("/") if p]

        # Upload endpoints: POST /api/upload/{target}
        if len(parts) == 3 and parts[:2] == ["api", "upload"]:
            try:
                handle_upload(self, parts[2])
            except Exception as e:
                import traceback
                self.json_response(500, {"error": f"Server fout: {e}", "trace": traceback.format_exc()})
            return

        # Temporary DB restore endpoint: POST /api/restore-db?secret=<RESTORE_SECRET>
        if parts == ["api", "restore-db"]:
            from urllib.parse import parse_qs
            secret = parse_qs(urlparse(self.path).query).get("secret", [""])[0]
            expected = os.environ.get("RESTORE_SECRET", "")
            if not expected or secret != expected:
                self.json_response(403, {"error": "Verboden."})
                return
            length = int(self.headers.get("Content-Length", "0"))
            db_bytes = self.rfile.read(length)
            DB_PATH.write_bytes(db_bytes)
            self.json_response(200, {"ok": True, "bytes": len(db_bytes)})
            return

        if parts == ["api", "export-ohw"]:
            try:
                body = self.read_json_body()
                rows = body.get("rows", [])
                entity_label = body.get("entityLabel", "RPL")
                if openpyxl is None:
                    self.json_response(500, {"error": "openpyxl niet beschikbaar."})
                    return
                import io
                from openpyxl.styles import Font, PatternFill, Alignment
                from openpyxl.utils import get_column_letter
                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = "OHW"
                header_font = Font(bold=True, color="FFFFFF")
                header_fill = PatternFill("solid", fgColor="1565C0")
                total_font = Font(bold=True)
                total_fill = PatternFill("solid", fgColor="DEEAFF")
                headers = [
                    "Project", "Projectleider", "OHW vorig jaar",
                    "Gepresteerde omzet", "Verstuurde facturen",
                    "Ontvangen facturen", "Huidig OHW"
                ]
                ws.append(headers)
                for cell in ws[1]:
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = Alignment(horizontal="center")
                totals = {k: 0.0 for k in ["openingOHW", "performedRevenue", "invoiced", "vendorInvoiced", "currentOHW"]}
                for p in rows:
                    title = f"{p.get('projectNumber','')} - {p.get('projectName','')}" if p.get("projectName") else p.get("projectNumber", "")
                    ws.append([
                        title,
                        p.get("projectLead") or "-",
                        round(float(p.get("openingOHW") or 0), 2),
                        round(float(p.get("performedRevenue") or 0), 2),
                        round(float(p.get("invoiced") or 0), 2),
                        round(float(p.get("vendorInvoiced") or 0), 2),
                        round(float(p.get("currentOHW") or 0), 2),
                    ])
                    for k, col in [("openingOHW", 3), ("performedRevenue", 4), ("invoiced", 5), ("vendorInvoiced", 6), ("currentOHW", 7)]:
                        totals[k] += float(p.get(k) or 0)
                        ws.cell(row=ws.max_row, column=col).number_format = '#,##0.00'
                total_row = ["Totaal", "-"] + [round(totals[k], 2) for k in ["openingOHW", "performedRevenue", "invoiced", "vendorInvoiced", "currentOHW"]]
                ws.append(total_row)
                for cell in ws[ws.max_row]:
                    cell.font = total_font
                    cell.fill = total_fill
                    if cell.column >= 3:
                        cell.number_format = '#,##0.00'
                col_widths = [60, 25, 18, 20, 20, 20, 15]
                for i, width in enumerate(col_widths, 1):
                    ws.column_dimensions[get_column_letter(i)].width = width
                buf = io.BytesIO()
                wb.save(buf)
                xlsx_bytes = buf.getvalue()
                self.send_response(200)
                self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                self.send_header("Content-Disposition", f'attachment; filename="ohw_export_{entity_label}.xlsx"')
                self.send_header("Content-Length", str(len(xlsx_bytes)))
                self.end_headers()
                self.wfile.write(xlsx_bytes)
            except Exception as e:
                self.json_response(500, {"error": str(e)})
            return

        if parts == ["api", "generate-pdf"]:
            try:
                length = int(self.headers.get("Content-Length", "0"))
                html_bytes = self.rfile.read(length)
                html_content = html_bytes.decode("utf-8")
                filename = self.headers.get("X-Filename", "rapport.pdf")

                chrome_paths = [
                    r"C:\Program Files\Google\Chrome\Application\chrome.exe",
                    r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe",
                    r"C:\Program Files\Chromium\Application\chrome.exe",
                ]
                chrome = next((p for p in chrome_paths if os.path.exists(p)), None)
                if not chrome:
                    self.json_response(500, {"error": "Chrome niet gevonden op dit systeem."})
                    return

                tmp_dir = tempfile.mkdtemp()
                try:
                    html_path = os.path.join(tmp_dir, "rapport.html")
                    pdf_path = os.path.join(tmp_dir, "rapport.pdf")
                    user_data_dir = os.path.join(tmp_dir, "chrome_profile")
                    with open(html_path, "w", encoding="utf-8") as f:
                        f.write(html_content)
                    subprocess.run([
                        chrome,
                        "--headless=new",
                        "--disable-gpu",
                        "--no-sandbox",
                        f"--user-data-dir={user_data_dir}",
                        f"--print-to-pdf={pdf_path}",
                        "--print-to-pdf-no-header",
                        html_path,
                    ], check=True, capture_output=True, timeout=30)
                    with open(pdf_path, "rb") as pf:
                        pdf_data = pf.read()
                finally:
                    import shutil
                    shutil.rmtree(tmp_dir, ignore_errors=True)

                self.send_response(200)
                self.send_header("Content-Type", "application/pdf")
                self.send_header("Content-Disposition", f'attachment; filename="{filename}"')
                self.send_header("Content-Length", str(len(pdf_data)))
                self.end_headers()
                self.wfile.write(pdf_data)
            except subprocess.TimeoutExpired:
                self.json_response(500, {"error": "PDF generatie time-out."})
            except Exception as e:
                self.json_response(500, {"error": str(e)})
            return

        try:
            payload = self.read_json_body()
        except Exception:
            self.json_response(400, {"error": "Ongeldige JSON."})
            return

        if parts == ["api", "consultants"]:
            name = (payload.get("name") or "").strip()
            start_date = payload.get("startDate") or ""
            exit_date = payload.get("exitDate") or ""
            entity = payload.get("entity") or "RPL BE"
            if entity not in {"RPL BE", "RPL NL"}:
                self.json_response(400, {"error": "Entiteit moet RPL BE of RPL NL zijn."})
                return
            if not name or parse_iso_date(start_date) is None:
                self.json_response(400, {"error": "Naam en geldige indiensttreding zijn verplicht."})
                return
            if exit_date and parse_iso_date(exit_date) is None:
                self.json_response(400, {"error": "Uit dienst moet een geldige datum zijn."})
                return

            with get_db() as conn:
                conn.execute(
                    "INSERT INTO consultants (name, start_date, entity, exit_date) VALUES (?, ?, ?, ?)",
                    (name, start_date, entity, exit_date or None),
                )
                data = list_consultants(conn)
            self.json_response(201, {"consultants": data})
            return

        if parts == ["api", "vacations"]:
            consultant_id = payload.get("consultantId")
            start_date = payload.get("startDate") or ""
            end_date = payload.get("endDate") or ""

            if not isinstance(consultant_id, int):
                self.json_response(400, {"error": "Consultant is verplicht."})
                return
            if parse_iso_date(start_date) is None or parse_iso_date(end_date) is None:
                self.json_response(400, {"error": "Start- en einddatum moeten geldig zijn."})
                return
            if end_date < start_date:
                self.json_response(400, {"error": "Einddatum kan niet voor startdatum liggen."})
                return

            today_iso = datetime.now().strftime("%Y-%m-%d")
            with get_db() as conn:
                consultant = conn.execute(
                    "SELECT id, start_date, exit_date FROM consultants WHERE id = ?",
                    (consultant_id,),
                ).fetchone()
                if consultant is None:
                    self.json_response(404, {"error": "Consultant niet gevonden."})
                    return
                if consultant["start_date"] > today_iso or (
                    consultant["exit_date"] and consultant["exit_date"] < today_iso
                ):
                    self.json_response(400, {"error": "Je kan enkel actieve consultants selecteren."})
                    return
                conn.execute(
                    """
                    INSERT INTO consultant_vacations (consultant_id, start_date, end_date)
                    VALUES (?, ?, ?)
                    """,
                    (consultant_id, start_date, end_date),
                )
                data = list_consultants(conn)
            self.json_response(201, {"consultants": data})
            return

        if len(parts) == 4 and parts[0] == "api" and parts[1] == "vacations" and parts[3] == "delete":
            try:
                vacation_id = int(parts[2])
            except Exception:
                self.json_response(400, {"error": "Ongeldige vakantie id."})
                return
            with get_db() as conn:
                row = conn.execute(
                    "SELECT id FROM consultant_vacations WHERE id = ?",
                    (vacation_id,),
                ).fetchone()
                if row is None:
                    self.json_response(404, {"error": "Vakantie niet gevonden."})
                    return
                conn.execute("DELETE FROM consultant_vacations WHERE id = ?", (vacation_id,))
                data = list_consultants(conn)
            self.json_response(200, {"consultants": data})
            return

        if len(parts) == 4 and parts[0] == "api" and parts[1] == "consultants" and parts[3] == "update":
            try:
                consultant_id = int(parts[2])
            except Exception:
                self.json_response(400, {"error": "Ongeldige consultant id."})
                return

            name = (payload.get("name") or "").strip()
            start_date = payload.get("startDate") or ""
            exit_date = payload.get("exitDate") or ""
            entity = payload.get("entity") or "RPL BE"
            if entity not in {"RPL BE", "RPL NL"}:
                self.json_response(400, {"error": "Entiteit moet RPL BE of RPL NL zijn."})
                return
            if not name or parse_iso_date(start_date) is None:
                self.json_response(400, {"error": "Naam en geldige indiensttreding zijn verplicht."})
                return
            if exit_date and parse_iso_date(exit_date) is None:
                self.json_response(400, {"error": "Uit dienst moet een geldige datum zijn."})
                return
            if exit_date and parse_iso_date(exit_date) < parse_iso_date(start_date):
                self.json_response(400, {"error": "Uit dienst kan niet voor indienst liggen."})
                return

            with get_db() as conn:
                consultant = conn.execute("SELECT id FROM consultants WHERE id = ?", (consultant_id,)).fetchone()
                if consultant is None:
                    self.json_response(404, {"error": "Consultant niet gevonden."})
                    return
                conn.execute(
                    "UPDATE consultants SET name = ?, start_date = ?, entity = ?, exit_date = ? WHERE id = ?",
                    (name, start_date, entity, exit_date or None, consultant_id),
                )
                data = list_consultants(conn)
            self.json_response(200, {"consultants": data})
            return

        if parts == ["api", "overhead-costs"]:
            year = payload.get("year")
            entity = (payload.get("entity") or "").strip()
            cost = payload.get("cost")
            if not isinstance(year, int) or year < 1900 or year > 3000:
                self.json_response(400, {"error": "Jaar is ongeldig."})
                return
            if entity not in {"RPL BE", "RPL NL"}:
                self.json_response(400, {"error": "Entiteit moet RPL BE of RPL NL zijn."})
                return
            try:
                cost_value = float(cost)
            except Exception:
                self.json_response(400, {"error": "Overheadkost moet numeriek zijn."})
                return
            if cost_value < 0:
                self.json_response(400, {"error": "Overheadkost kan niet negatief zijn."})
                return
            with get_db() as conn:
                conn.execute(
                    """
                    INSERT INTO overhead_cost_history (year, entity, cost)
                    VALUES (?, ?, ?)
                    ON CONFLICT(year, entity) DO UPDATE SET cost = excluded.cost
                    """,
                    (year, entity, cost_value),
                )
                data = list_overhead_costs(conn)
            self.json_response(201, {"overheadCosts": data})
            return

        if len(parts) == 4 and parts[0] == "api" and parts[1] == "overhead-costs" and parts[3] == "delete":
            try:
                overhead_id = int(parts[2])
            except Exception:
                self.json_response(400, {"error": "Ongeldige overhead id."})
                return
            with get_db() as conn:
                row = conn.execute("SELECT id FROM overhead_cost_history WHERE id = ?", (overhead_id,)).fetchone()
                if row is None:
                    self.json_response(404, {"error": "Overheadkost niet gevonden."})
                    return
                conn.execute("DELETE FROM overhead_cost_history WHERE id = ?", (overhead_id,))
                data = list_overhead_costs(conn)
            self.json_response(200, {"overheadCosts": data})
            return

        if parts == ["api", "offers"]:
            name = (payload.get("name") or "").strip()
            submission_date = payload.get("submissionDate") or ""
            if not name or parse_iso_date(submission_date) is None:
                self.json_response(400, {"error": "Naam en geldige indieningsdatum zijn verplicht."})
                return
            with get_db() as conn:
                conn.execute(
                    "INSERT INTO offers (name, created_date, submission_date, status, close_reason) VALUES (?, date('now'), ?, 'open', NULL)",
                    (name, submission_date),
                )
                data = list_offers(conn)
            self.json_response(201, {"offers": data})
            return

        if parts == ["api", "projects"]:
            project_number = (payload.get("projectNumber") or "").strip()
            name = (payload.get("name") or "").strip()
            start_date = payload.get("startDate") or ""
            delivery_date = payload.get("deliveryDate") or ""
            budget = payload.get("budget")
            status_label = payload.get("status") or "actief"
            lead_consultant_id = payload.get("leadConsultantId")
            billing_type = (payload.get("billingType") or "regie").strip().lower()
            previous_budget_total_raw = payload.get("previousBudgetTotal")
            next_budget_total_raw = payload.get("nextBudgetTotal")
            cap_amount_raw = payload.get("capAmount")
            if not project_number or not name or parse_iso_date(delivery_date) is None:
                self.json_response(400, {"error": "Projectnummer, naam en geldige opleveringsdatum zijn verplicht."})
                return
            if start_date and parse_iso_date(start_date) is None:
                self.json_response(400, {"error": "Startdatum moet een geldige datum zijn."})
                return
            effective_start = start_date or datetime.now().strftime("%Y-%m-%d")
            if parse_iso_date(effective_start) > parse_iso_date(delivery_date):
                self.json_response(400, {"error": "Startdatum kan niet na opleveringsdatum liggen."})
                return
            if billing_type not in {"fixed", "regie", "regie_cap"}:
                self.json_response(400, {"error": "Type project moet fixed, regie of regie met cap zijn."})
                return
            if not isinstance(lead_consultant_id, int):
                self.json_response(400, {"error": "Projectleider is ongeldig."})
                return
            try:
                budget_value = float(budget)
            except Exception:
                self.json_response(400, {"error": "Budget moet numeriek zijn."})
                return
            try:
                previous_budget_total = parse_optional_budget_total(previous_budget_total_raw)
                next_budget_total = parse_optional_budget_total(next_budget_total_raw)
                cap_amount = parse_optional_budget_total(cap_amount_raw)
            except ValueError as ex:
                self.json_response(400, {"error": str(ex)})
                return
            if billing_type == "regie":
                budget_value = 0.0
                previous_budget_total = 0.0
                next_budget_total = 0.0
                cap_amount = 0.0
            elif billing_type == "regie_cap":
                if cap_amount <= 0:
                    self.json_response(400, {"error": "Cap moet groter dan 0 zijn voor regie met cap."})
                    return
                budget_value = 0.0
                if previous_budget_total + next_budget_total > cap_amount:
                    self.json_response(400, {"error": "Som van voorgaande en komende jaren kan niet groter zijn dan de cap."})
                    return
            elif budget_value < 0:
                self.json_response(400, {"error": "Budget kan niet negatief zijn."})
                return
            if status_label not in {"actief", "on hold", "afgerond", "geannuleerd"}:
                self.json_response(400, {"error": "Ongeldige status."})
                return
            with get_db() as conn:
                lead = conn.execute(
                    "SELECT id, start_date, exit_date FROM consultants WHERE id = ?",
                    (lead_consultant_id,),
                ).fetchone()
                if lead is None:
                    self.json_response(400, {"error": "Projectleider niet gevonden."})
                    return
                today_iso = datetime.now().strftime("%Y-%m-%d")
                if lead["start_date"] > today_iso or (lead["exit_date"] and lead["exit_date"] < today_iso):
                    self.json_response(400, {"error": "Projectleider moet een actieve consultant zijn."})
                    return
                conn.execute(
                    """
                    INSERT INTO projects
                    (project_number, name, created_date, start_date, delivery_date, budget, status, lead_consultant_id, billing_type, previous_budget_total, next_budget_total, cap_amount)
                    VALUES (?, ?, date('now'), ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        project_number,
                        name,
                        effective_start,
                        delivery_date,
                        budget_value,
                        status_label,
                        lead_consultant_id,
                        billing_type,
                        previous_budget_total,
                        next_budget_total,
                        cap_amount,
                    ),
                )
                data = list_projects(conn)
            self.json_response(201, {"projects": data})
            return

        if len(parts) == 4 and parts[0] == "api" and parts[1] == "offers" and parts[3] == "staff":
            try:
                offer_id = int(parts[2])
            except Exception:
                self.json_response(400, {"error": "Ongeldige offerte id."})
                return
            consultant_id = payload.get("consultantId")
            if not isinstance(consultant_id, int):
                self.json_response(400, {"error": "Consultant is verplicht."})
                return
            with get_db() as conn:
                offer = conn.execute("SELECT id, status FROM offers WHERE id = ?", (offer_id,)).fetchone()
                if offer is None:
                    self.json_response(404, {"error": "Offerte niet gevonden."})
                    return
                if offer["status"] != "open":
                    self.json_response(400, {"error": "Offerte is al afgesloten."})
                    return
                consultant = conn.execute("SELECT id FROM consultants WHERE id = ?", (consultant_id,)).fetchone()
                if consultant is None:
                    self.json_response(400, {"error": "Consultant niet gevonden."})
                    return
                conn.execute(
                    "INSERT OR IGNORE INTO offer_staff (offer_id, consultant_id) VALUES (?, ?)",
                    (offer_id, consultant_id),
                )
                data = list_offers(conn)
            self.json_response(200, {"offers": data})
            return

        if len(parts) == 4 and parts[0] == "api" and parts[1] == "offers" and parts[3] == "allocations":
            try:
                offer_id = int(parts[2])
            except Exception:
                self.json_response(400, {"error": "Ongeldige offerte id."})
                return

            consultant_id = payload.get("consultantId")
            week_start = payload.get("weekStart") or ""
            mandays = payload.get("mandays")

            if not isinstance(consultant_id, int):
                self.json_response(400, {"error": "Consultant is verplicht."})
                return
            if parse_iso_date(week_start) is None:
                self.json_response(400, {"error": "Weekstart moet yyyy-mm-dd zijn."})
                return
            try:
                mandays_value = float(mandays)
            except Exception:
                self.json_response(400, {"error": "Mandagen moeten numeriek zijn."})
                return
            if mandays_value < 0:
                self.json_response(400, {"error": "Mandagen kunnen niet negatief zijn."})
                return

            with get_db() as conn:
                offer = conn.execute(
                    "SELECT id, submission_date, status FROM offers WHERE id = ?",
                    (offer_id,),
                ).fetchone()
                if offer is None:
                    self.json_response(404, {"error": "Offerte niet gevonden."})
                    return
                if offer["status"] != "open":
                    self.json_response(400, {"error": "Offerte is al afgesloten."})
                    return
                if parse_iso_date(week_start) > parse_iso_date(offer["submission_date"]):
                    self.json_response(400, {"error": "Weekstart kan niet na indieningsdatum liggen."})
                    return
                consultant = conn.execute("SELECT id FROM consultants WHERE id = ?", (consultant_id,)).fetchone()
                if consultant is None:
                    self.json_response(400, {"error": "Consultant niet gevonden."})
                    return

                conn.execute(
                    "INSERT OR IGNORE INTO offer_staff (offer_id, consultant_id) VALUES (?, ?)",
                    (offer_id, consultant_id),
                )
                existing = conn.execute(
                    """
                    SELECT id FROM offer_allocations
                    WHERE offer_id = ? AND consultant_id = ? AND week_start = ?
                    """,
                    (offer_id, consultant_id, week_start),
                ).fetchone()
                if mandays_value == 0:
                    if existing:
                        conn.execute("DELETE FROM offer_allocations WHERE id = ?", (existing["id"],))
                elif existing:
                    conn.execute("UPDATE offer_allocations SET mandays = ? WHERE id = ?", (mandays_value, existing["id"]))
                else:
                    conn.execute(
                        """
                        INSERT INTO offer_allocations (offer_id, consultant_id, week_start, mandays)
                        VALUES (?, ?, ?, ?)
                        """,
                        (offer_id, consultant_id, week_start, mandays_value),
                    )
                data = list_offers(conn)
            self.json_response(201, {"offers": data})
            return

        if len(parts) == 4 and parts[0] == "api" and parts[1] == "offers" and parts[3] == "close":
            try:
                offer_id = int(parts[2])
            except Exception:
                self.json_response(400, {"error": "Ongeldige offerte id."})
                return

            reason = payload.get("reason") or ""
            if reason not in {"gewonnen", "verloren", "dropped", "geannuleerd"}:
                self.json_response(400, {"error": "Ongeldige afsluitreden."})
                return

            with get_db() as conn:
                offer = conn.execute("SELECT id, status FROM offers WHERE id = ?", (offer_id,)).fetchone()
                if offer is None:
                    self.json_response(404, {"error": "Offerte niet gevonden."})
                    return
                if offer["status"] != "open":
                    self.json_response(400, {"error": "Offerte is al afgesloten."})
                    return
                conn.execute(
                    "UPDATE offers SET status = 'closed', close_reason = ? WHERE id = ?",
                    (reason, offer_id),
                )
                data = list_offers(conn)
            self.json_response(200, {"offers": data})
            return

        if len(parts) == 4 and parts[0] == "api" and parts[1] == "offers" and parts[3] == "update":
            try:
                offer_id = int(parts[2])
            except Exception:
                self.json_response(400, {"error": "Ongeldige offerte id."})
                return

            name = (payload.get("name") or "").strip()
            submission_date = payload.get("submissionDate") or ""
            status_label = payload.get("status") or "open"
            if not name or parse_iso_date(submission_date) is None:
                self.json_response(400, {"error": "Titel en geldige indieningsdatum zijn verplicht."})
                return
            if status_label not in {"open", "gewonnen", "verloren", "dropped", "geannuleerd"}:
                self.json_response(400, {"error": "Ongeldige status."})
                return

            with get_db() as conn:
                offer = conn.execute("SELECT id FROM offers WHERE id = ?", (offer_id,)).fetchone()
                if offer is None:
                    self.json_response(404, {"error": "Offerte niet gevonden."})
                    return
                if status_label == "open":
                    conn.execute(
                        "UPDATE offers SET name = ?, submission_date = ?, status = 'open', close_reason = NULL WHERE id = ?",
                        (name, submission_date, offer_id),
                    )
                else:
                    conn.execute(
                        "UPDATE offers SET name = ?, submission_date = ?, status = 'closed', close_reason = ? WHERE id = ?",
                        (name, submission_date, status_label, offer_id),
                    )
                data = list_offers(conn)
            self.json_response(200, {"offers": data})
            return

        if len(parts) == 4 and parts[0] == "api" and parts[1] == "offers" and parts[3] == "delete":
            try:
                offer_id = int(parts[2])
            except Exception:
                self.json_response(400, {"error": "Ongeldige offerte id."})
                return
            with get_db() as conn:
                offer = conn.execute("SELECT id FROM offers WHERE id = ?", (offer_id,)).fetchone()
                if offer is None:
                    self.json_response(404, {"error": "Offerte niet gevonden."})
                    return
                conn.execute("DELETE FROM offers WHERE id = ?", (offer_id,))
                data = list_offers(conn)
            self.json_response(200, {"offers": data})
            return

        if len(parts) == 4 and parts[0] == "api" and parts[1] == "projects" and parts[3] == "update":
            try:
                project_id = int(parts[2])
            except Exception:
                self.json_response(400, {"error": "Ongeldige project id."})
                return

            project_number = (payload.get("projectNumber") or "").strip()
            name = (payload.get("name") or "").strip()
            start_date = payload.get("startDate") or ""
            delivery_date = payload.get("deliveryDate") or ""
            budget = payload.get("budget")
            status_label = payload.get("status") or "actief"
            lead_consultant_id = payload.get("leadConsultantId")
            billing_type = (payload.get("billingType") or "regie").strip().lower()
            previous_budget_total_raw = payload.get("previousBudgetTotal")
            next_budget_total_raw = payload.get("nextBudgetTotal")
            cap_amount_raw = payload.get("capAmount")
            if not project_number or not name or parse_iso_date(delivery_date) is None:
                self.json_response(400, {"error": "Projectnummer, titel en geldige opleveringsdatum zijn verplicht."})
                return
            if start_date and parse_iso_date(start_date) is None:
                self.json_response(400, {"error": "Startdatum moet een geldige datum zijn."})
                return
            if billing_type not in {"fixed", "regie", "regie_cap"}:
                self.json_response(400, {"error": "Type project moet fixed, regie of regie met cap zijn."})
                return
            if not isinstance(lead_consultant_id, int):
                self.json_response(400, {"error": "Projectleider is ongeldig."})
                return
            try:
                budget_value = float(budget)
            except Exception:
                self.json_response(400, {"error": "Budget moet numeriek zijn."})
                return
            try:
                previous_budget_total = parse_optional_budget_total(previous_budget_total_raw)
                next_budget_total = parse_optional_budget_total(next_budget_total_raw)
                cap_amount = parse_optional_budget_total(cap_amount_raw)
            except ValueError as ex:
                self.json_response(400, {"error": str(ex)})
                return
            if billing_type == "regie":
                budget_value = 0.0
                previous_budget_total = 0.0
                next_budget_total = 0.0
                cap_amount = 0.0
            elif billing_type == "regie_cap":
                if cap_amount <= 0:
                    self.json_response(400, {"error": "Cap moet groter dan 0 zijn voor regie met cap."})
                    return
                budget_value = 0.0
                if previous_budget_total + next_budget_total > cap_amount:
                    self.json_response(400, {"error": "Som van voorgaande en komende jaren kan niet groter zijn dan de cap."})
                    return
            elif budget_value < 0:
                self.json_response(400, {"error": "Budget kan niet negatief zijn."})
                return
            if status_label not in {"actief", "on hold", "afgerond", "geannuleerd"}:
                self.json_response(400, {"error": "Ongeldige status."})
                return

            with get_db() as conn:
                project = conn.execute("SELECT id, created_date FROM projects WHERE id = ?", (project_id,)).fetchone()
                if project is None:
                    self.json_response(404, {"error": "Project niet gevonden."})
                    return
                lead = conn.execute(
                    "SELECT id, start_date, exit_date FROM consultants WHERE id = ?",
                    (lead_consultant_id,),
                ).fetchone()
                if lead is None:
                    self.json_response(400, {"error": "Projectleider niet gevonden."})
                    return
                today_iso = datetime.now().strftime("%Y-%m-%d")
                if lead["start_date"] > today_iso or (lead["exit_date"] and lead["exit_date"] < today_iso):
                    self.json_response(400, {"error": "Projectleider moet een actieve consultant zijn."})
                    return
                effective_start = start_date or project["created_date"]
                if parse_iso_date(effective_start) > parse_iso_date(delivery_date):
                    self.json_response(400, {"error": "Startdatum kan niet na opleveringsdatum liggen."})
                    return
                conn.execute(
                    """
                    UPDATE projects
                    SET project_number = ?, name = ?, start_date = ?, delivery_date = ?, budget = ?, status = ?, lead_consultant_id = ?, billing_type = ?, previous_budget_total = ?, next_budget_total = ?, cap_amount = ?
                    WHERE id = ?
                    """,
                    (
                        project_number,
                        name,
                        effective_start,
                        delivery_date,
                        budget_value,
                        status_label,
                        lead_consultant_id,
                        billing_type,
                        previous_budget_total,
                        next_budget_total,
                        cap_amount,
                        project_id,
                    ),
                )
                data = list_projects(conn)
            self.json_response(200, {"projects": data})
            return

        if len(parts) == 4 and parts[0] == "api" and parts[1] == "projects" and parts[3] == "staff":
            try:
                project_id = int(parts[2])
            except Exception:
                self.json_response(400, {"error": "Ongeldige project id."})
                return
            consultant_id = payload.get("consultantId")
            hourly_rate = payload.get("hourlyRate", 0)
            if not isinstance(consultant_id, int):
                self.json_response(400, {"error": "Consultant is verplicht."})
                return
            try:
                hourly_rate_value = float(hourly_rate)
            except Exception:
                self.json_response(400, {"error": "Tarief per uur moet numeriek zijn."})
                return
            if hourly_rate_value < 0:
                self.json_response(400, {"error": "Tarief per uur kan niet negatief zijn."})
                return
            with get_db() as conn:
                project = conn.execute("SELECT id, status FROM projects WHERE id = ?", (project_id,)).fetchone()
                if project is None:
                    self.json_response(404, {"error": "Project niet gevonden."})
                    return
                if project["status"] != "actief":
                    self.json_response(400, {"error": "Project is niet actief."})
                    return
                consultant = conn.execute("SELECT id FROM consultants WHERE id = ?", (consultant_id,)).fetchone()
                if consultant is None:
                    self.json_response(400, {"error": "Consultant niet gevonden."})
                    return
                conn.execute(
                    """
                    INSERT INTO project_staff (project_id, consultant_id, hourly_rate)
                    VALUES (?, ?, ?)
                    ON CONFLICT(project_id, consultant_id) DO UPDATE SET hourly_rate = excluded.hourly_rate
                    """,
                    (project_id, consultant_id, hourly_rate_value),
                )
                data = list_projects(conn)
            self.json_response(200, {"projects": data})
            return

        if len(parts) == 4 and parts[0] == "api" and parts[1] == "projects" and parts[3] == "staff-rate":
            try:
                project_id = int(parts[2])
            except Exception:
                self.json_response(400, {"error": "Ongeldige project id."})
                return
            consultant_id = payload.get("consultantId")
            hourly_rate = payload.get("hourlyRate")
            if not isinstance(consultant_id, int):
                self.json_response(400, {"error": "Consultant is verplicht."})
                return
            try:
                hourly_rate_value = float(hourly_rate)
            except Exception:
                self.json_response(400, {"error": "Tarief per uur moet numeriek zijn."})
                return
            if hourly_rate_value < 0:
                self.json_response(400, {"error": "Tarief per uur kan niet negatief zijn."})
                return
            with get_db() as conn:
                staff = conn.execute(
                    "SELECT id FROM project_staff WHERE project_id = ? AND consultant_id = ?",
                    (project_id, consultant_id),
                ).fetchone()
                if staff is None:
                    self.json_response(404, {"error": "Consultant is nog niet toegevoegd aan dit project."})
                    return
                conn.execute(
                    "UPDATE project_staff SET hourly_rate = ? WHERE project_id = ? AND consultant_id = ?",
                    (hourly_rate_value, project_id, consultant_id),
                )
                data = list_projects(conn)
            self.json_response(200, {"projects": data})
            return

        if len(parts) == 4 and parts[0] == "api" and parts[1] == "projects" and parts[3] == "staff-delete":
            try:
                project_id = int(parts[2])
            except Exception:
                self.json_response(400, {"error": "Ongeldige project id."})
                return
            consultant_id = payload.get("consultantId")
            if not isinstance(consultant_id, int):
                self.json_response(400, {"error": "Consultant is verplicht."})
                return
            with get_db() as conn:
                project = conn.execute("SELECT id FROM projects WHERE id = ?", (project_id,)).fetchone()
                if project is None:
                    self.json_response(404, {"error": "Project niet gevonden."})
                    return
                conn.execute(
                    "DELETE FROM project_allocations WHERE project_id = ? AND consultant_id = ?",
                    (project_id, consultant_id),
                )
                conn.execute(
                    "DELETE FROM project_staff WHERE project_id = ? AND consultant_id = ?",
                    (project_id, consultant_id),
                )
                data = list_projects(conn)
            self.json_response(200, {"projects": data})
            return

        if len(parts) == 4 and parts[0] == "api" and parts[1] == "projects" and parts[3] == "allocations":
            try:
                project_id = int(parts[2])
            except Exception:
                self.json_response(400, {"error": "Ongeldige project id."})
                return

            consultant_id = payload.get("consultantId")
            week_start = payload.get("weekStart") or ""
            mandays = payload.get("mandays")
            if not isinstance(consultant_id, int):
                self.json_response(400, {"error": "Consultant is verplicht."})
                return
            if parse_iso_date(week_start) is None:
                self.json_response(400, {"error": "Weekstart moet yyyy-mm-dd zijn."})
                return
            try:
                mandays_value = float(mandays)
            except Exception:
                self.json_response(400, {"error": "Mandagen moeten numeriek zijn."})
                return
            if mandays_value < 0:
                self.json_response(400, {"error": "Mandagen kunnen niet negatief zijn."})
                return

            with get_db() as conn:
                project = conn.execute(
                    "SELECT id, start_date, created_date, delivery_date, status FROM projects WHERE id = ?",
                    (project_id,),
                ).fetchone()
                if project is None:
                    self.json_response(404, {"error": "Project niet gevonden."})
                    return
                if project["status"] != "actief":
                    self.json_response(400, {"error": "Project is niet actief."})
                    return
                effective_start = project["start_date"] or project["created_date"]
                if parse_iso_date(week_start) < parse_iso_date(effective_start):
                    self.json_response(400, {"error": "Weekstart kan niet voor startdatum liggen."})
                    return
                if parse_iso_date(week_start) > parse_iso_date(project["delivery_date"]):
                    self.json_response(400, {"error": "Weekstart kan niet na opleveringsdatum liggen."})
                    return
                if week_fully_passed(week_start):
                    self.json_response(400, {"error": "Voor volledig voorbije weken wordt data uit timesheets gebruikt en kan je niet manueel wijzigen."})
                    return
                consultant = conn.execute("SELECT id FROM consultants WHERE id = ?", (consultant_id,)).fetchone()
                if consultant is None:
                    self.json_response(400, {"error": "Consultant niet gevonden."})
                    return

                conn.execute(
                    "INSERT OR IGNORE INTO project_staff (project_id, consultant_id) VALUES (?, ?)",
                    (project_id, consultant_id),
                )
                existing = conn.execute(
                    """
                    SELECT id FROM project_allocations
                    WHERE project_id = ? AND consultant_id = ? AND week_start = ?
                    """,
                    (project_id, consultant_id, week_start),
                ).fetchone()
                if mandays_value == 0:
                    if existing:
                        conn.execute("DELETE FROM project_allocations WHERE id = ?", (existing["id"],))
                elif existing:
                    conn.execute("UPDATE project_allocations SET mandays = ? WHERE id = ?", (mandays_value, existing["id"]))
                else:
                    conn.execute(
                        """
                        INSERT INTO project_allocations (project_id, consultant_id, week_start, mandays)
                        VALUES (?, ?, ?, ?)
                        """,
                        (project_id, consultant_id, week_start, mandays_value),
                    )
                data = list_projects(conn)
            self.json_response(201, {"projects": data})
            return

        if len(parts) == 4 and parts[0] == "api" and parts[1] == "projects" and parts[3] == "delete":
            try:
                project_id = int(parts[2])
            except Exception:
                self.json_response(400, {"error": "Ongeldige project id."})
                return
            with get_db() as conn:
                project = conn.execute("SELECT id FROM projects WHERE id = ?", (project_id,)).fetchone()
                if project is None:
                    self.json_response(404, {"error": "Project niet gevonden."})
                    return
                conn.execute("DELETE FROM projects WHERE id = ?", (project_id,))
                data = list_projects(conn)
            self.json_response(200, {"projects": data})
            return

        if len(parts) == 4 and parts[0] == "api" and parts[1] == "consultants" and parts[3] in ("role", "regime"):
            try:
                consultant_id = int(parts[2])
            except Exception:
                self.json_response(400, {"error": "Ongeldige consultant id."})
                return

            with get_db() as conn:
                if parts[3] == "role":
                    role = payload.get("role") or ""
                    from_date = payload.get("from") or ""
                    to_date = payload.get("to") or ""
                    if role not in {"junior", "medior", "senior", "expert", "director"}:
                        self.json_response(400, {"error": "Ongeldige functie."})
                        return
                    error = validate_timeline(conn, "role_history", consultant_id, from_date, to_date)
                    if error:
                        self.json_response(400, {"error": error})
                        return
                    conn.execute(
                        "INSERT INTO role_history (consultant_id, role, from_date, to_date) VALUES (?, ?, ?, ?)",
                        (consultant_id, role, from_date, to_date or None),
                    )
                else:
                    regime = payload.get("regime")
                    from_date = payload.get("from") or ""
                    to_date = payload.get("to") or ""
                    if not isinstance(regime, int) or regime < 1 or regime > 100:
                        self.json_response(400, {"error": "Werkregime moet tussen 1 en 100 liggen."})
                        return
                    error = validate_timeline(conn, "regime_history", consultant_id, from_date, to_date)
                    if error:
                        self.json_response(400, {"error": error})
                        return
                    conn.execute(
                        "INSERT INTO regime_history (consultant_id, regime, from_date, to_date) VALUES (?, ?, ?, ?)",
                        (consultant_id, regime, from_date, to_date or None),
                    )

                data = list_consultants(conn)
            self.json_response(201, {"consultants": data})
            return

        if len(parts) == 4 and parts[0] == "api" and parts[1] == "consultants" and parts[3] == "cost":
            try:
                consultant_id = int(parts[2])
            except Exception:
                self.json_response(400, {"error": "Ongeldige consultant id."})
                return
            year = payload.get("year")
            cost = payload.get("cost")
            if not isinstance(year, int):
                self.json_response(400, {"error": "Jaar is verplicht."})
                return
            if year < 1900 or year > 3000:
                self.json_response(400, {"error": "Jaar is ongeldig."})
                return
            try:
                cost_value = float(cost)
            except Exception:
                self.json_response(400, {"error": "Kostprijs moet numeriek zijn."})
                return
            if cost_value < 0:
                self.json_response(400, {"error": "Kostprijs kan niet negatief zijn."})
                return
            with get_db() as conn:
                consultant = conn.execute("SELECT id FROM consultants WHERE id = ?", (consultant_id,)).fetchone()
                if consultant is None:
                    self.json_response(404, {"error": "Consultant niet gevonden."})
                    return
                conn.execute(
                    """
                    INSERT INTO consultant_cost_history (consultant_id, year, cost)
                    VALUES (?, ?, ?)
                    ON CONFLICT(consultant_id, year) DO UPDATE SET cost = excluded.cost
                    """,
                    (consultant_id, year, cost_value),
                )
                data = list_consultants(conn)
            self.json_response(201, {"consultants": data})
            return

        if (
            len(parts) == 6
            and parts[0] == "api"
            and parts[1] == "consultants"
            and parts[3] in ("role", "regime", "cost")
            and parts[5] in ("update", "delete")
        ):
            try:
                consultant_id = int(parts[2])
                entry_id = int(parts[4])
            except Exception:
                self.json_response(400, {"error": "Ongeldige id."})
                return

            with get_db() as conn:
                if parts[3] == "role":
                    table = "role_history"
                elif parts[3] == "regime":
                    table = "regime_history"
                else:
                    table = "consultant_cost_history"
                row = conn.execute(
                    f"SELECT id FROM {table} WHERE id = ? AND consultant_id = ?",
                    (entry_id, consultant_id),
                ).fetchone()
                if row is None:
                    self.json_response(404, {"error": "Historiekregel niet gevonden."})
                    return

                if parts[5] == "delete":
                    conn.execute(
                        f"DELETE FROM {table} WHERE id = ? AND consultant_id = ?",
                        (entry_id, consultant_id),
                    )
                    data = list_consultants(conn)
                    self.json_response(200, {"consultants": data})
                    return

                if parts[3] == "role":
                    from_date = payload.get("from") or ""
                    to_date = payload.get("to") or ""
                    role = payload.get("role") or ""
                    if role not in {"junior", "medior", "senior", "expert", "director"}:
                        self.json_response(400, {"error": "Ongeldige functie."})
                        return
                    error = validate_timeline(conn, table, consultant_id, from_date, to_date, entry_id)
                    if error:
                        self.json_response(400, {"error": error})
                        return
                    conn.execute(
                        "UPDATE role_history SET role = ?, from_date = ?, to_date = ? WHERE id = ? AND consultant_id = ?",
                        (role, from_date, to_date or None, entry_id, consultant_id),
                    )
                elif parts[3] == "regime":
                    from_date = payload.get("from") or ""
                    to_date = payload.get("to") or ""
                    regime = payload.get("regime")
                    if not isinstance(regime, int) or regime < 1 or regime > 100:
                        self.json_response(400, {"error": "Werkregime moet tussen 1 en 100 liggen."})
                        return
                    error = validate_timeline(conn, table, consultant_id, from_date, to_date, entry_id)
                    if error:
                        self.json_response(400, {"error": error})
                        return
                    conn.execute(
                        "UPDATE regime_history SET regime = ?, from_date = ?, to_date = ? WHERE id = ? AND consultant_id = ?",
                        (regime, from_date, to_date or None, entry_id, consultant_id),
                    )
                else:
                    year = payload.get("year")
                    cost = payload.get("cost")
                    if not isinstance(year, int) or year < 1900 or year > 3000:
                        self.json_response(400, {"error": "Jaar is ongeldig."})
                        return
                    try:
                        cost_value = float(cost)
                    except Exception:
                        self.json_response(400, {"error": "Kostprijs moet numeriek zijn."})
                        return
                    if cost_value < 0:
                        self.json_response(400, {"error": "Kostprijs kan niet negatief zijn."})
                        return
                    existing = conn.execute(
                        """
                        SELECT id FROM consultant_cost_history
                        WHERE consultant_id = ? AND year = ? AND id != ?
                        """,
                        (consultant_id, year, entry_id),
                    ).fetchone()
                    if existing is not None:
                        self.json_response(400, {"error": "Voor dit jaar bestaat al een kostprijs."})
                        return
                    conn.execute(
                        "UPDATE consultant_cost_history SET year = ?, cost = ? WHERE id = ? AND consultant_id = ?",
                        (year, cost_value, entry_id, consultant_id),
                    )

                data = list_consultants(conn)
            self.json_response(200, {"consultants": data})
            return

        if len(parts) == 4 and parts[0] == "api" and parts[1] == "consultants" and parts[3] == "delete":
            try:
                consultant_id = int(parts[2])
            except Exception:
                self.json_response(400, {"error": "Ongeldige consultant id."})
                return
            with get_db() as conn:
                consultant = conn.execute("SELECT id FROM consultants WHERE id = ?", (consultant_id,)).fetchone()
                if consultant is None:
                    self.json_response(404, {"error": "Consultant niet gevonden."})
                    return
                conn.execute("DELETE FROM consultants WHERE id = ?", (consultant_id,))
                data = list_consultants(conn)
            self.json_response(200, {"consultants": data})
            return

        if parts == ["api", "export", "consultant"]:
            if openpyxl is None:
                self.json_response(500, {"error": "openpyxl is niet beschikbaar op de server."})
                return
            consultant_id = payload.get("consultantId")
            if not isinstance(consultant_id, int):
                self.json_response(400, {"error": "consultantId is verplicht."})
                return
            capacity_b64 = payload.get("capacityChartBase64", "")
            timeline_b64 = payload.get("timelineChartBase64", "")
            with get_db() as conn:
                file_data, result = export_consultant_planning(conn, consultant_id, capacity_b64, timeline_b64)
            if file_data is None:
                self.json_response(404, {"error": result})
                return
            safe_name = "".join(c for c in result if c.isalnum() or c in " _-").strip().replace(" ", "_")
            date_str = datetime.now().strftime("%Y%m%d")
            filename = f"planning_{safe_name}_{date_str}.xlsx"
            self.send_response(200)
            self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            self.send_header("Content-Disposition", f'attachment; filename="{filename}"')
            self.send_header("Content-Length", str(len(file_data)))
            self.end_headers()
            self.wfile.write(file_data)
            return


        self.json_response(404, {"error": "Endpoint niet gevonden."})


# ── Upload helpers ────────────────────────────────────────────────────────────

UPLOAD_TARGETS = {
    "timesheets-be": {
        "filename": "Employee Timesheets BE.xlsx",
        "sheet": "Exported Data",
        "key_cols": ["Employee Name", "Entry Date", "Project Name", "Task"],
        "label": "Employee Timesheets BE",
    },
    "timesheets-nl": {
        "filename": "Employee Timesheets NL.xlsx",
        "sheet": "Exported Data",
        "key_cols": ["Employee Name", "Entry Date", "Project Name", "Task"],
        "label": "Employee Timesheets NL",
    },
    "invoices": {
        "filename": "List of Invoices.xlsx",
        "sheet": None,  # first sheet
        "key_cols": ["Project No.", "Sales"],
        "label": "List of Invoices",
    },
    "projects": {
        "filename": "List of Projects.xlsx",
        "sheet": None,
        "key_cols": ["Project No."],
        "label": "List of Projects",
    },
    "vendor-invoices": {
        "filename": "Vendor Invoices.xlsx",
        "sheet": None,
        "key_cols": ["Project No"],
        "label": "Vendor Invoices",
    },
}


def _read_xlsx_rows(file_bytes, sheet_name=None):
    """Read all rows from an xlsx file (bytes). Returns (headers, list_of_dicts)."""
    if openpyxl is None:
        raise RuntimeError("openpyxl niet beschikbaar")
    wb = openpyxl.load_workbook(filename=io.BytesIO(file_bytes), read_only=True, data_only=True)
    if sheet_name and sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    else:
        ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []
    headers = [str(h).strip() if h is not None else f"col_{i}" for i, h in enumerate(rows[0])]
    data = []
    for row in rows[1:]:
        if all(v is None for v in row):
            continue
        data.append({headers[i]: (str(v).strip() if v is not None else "") for i, v in enumerate(row)})
    return headers, data


def _row_key(row, key_cols):
    """Build a hashable key from a row dict using the specified columns."""
    return tuple(row.get(c, "") for c in key_cols)


def _diff_xlsx(old_bytes, new_bytes, sheet_name, key_cols):
    """Compare two xlsx files and return added/removed rows."""
    _, old_rows = _read_xlsx_rows(old_bytes, sheet_name) if old_bytes else ([], [])
    headers, new_rows = _read_xlsx_rows(new_bytes, sheet_name)

    old_keys = {_row_key(r, key_cols): r for r in old_rows}
    new_keys = {_row_key(r, key_cols): r for r in new_rows}

    added = [new_keys[k] for k in new_keys if k not in old_keys]
    removed = [old_keys[k] for k in old_keys if k not in new_keys]
    unchanged = len([k for k in new_keys if k in old_keys])

    return headers, added, removed, unchanged


def handle_upload(handler, target_key):
    """Handle multipart file upload, return diff preview or save file."""
    cfg = UPLOAD_TARGETS.get(target_key)
    if cfg is None:
        handler.json_response(404, {"error": "Onbekend upload type."})
        return

    content_type = handler.headers.get("Content-Type", "")
    if "multipart/form-data" not in content_type:
        handler.json_response(400, {"error": "Verwacht multipart/form-data."})
        return

    if openpyxl is None:
        handler.json_response(500, {"error": "openpyxl niet beschikbaar op de server."})
        return

    # Parse multipart form data (cgi module removed in Python 3.13)
    content_length = int(handler.headers.get("Content-Length", "0"))
    raw_body = handler.rfile.read(content_length)

    # Extract boundary from Content-Type header
    boundary = None
    for ct_part in content_type.split(";"):
        ct_part = ct_part.strip()
        if ct_part.lower().startswith("boundary="):
            boundary = ct_part[9:].strip('"')
            break

    fields = {}
    if boundary:
        delimiter = ("--" + boundary).encode()
        for part in raw_body.split(delimiter)[1:]:
            # Stop at closing boundary
            if part.lstrip(b"\r\n").startswith(b"--"):
                break
            if b"\r\n\r\n" not in part:
                continue
            raw_headers, body = part.split(b"\r\n\r\n", 1)
            # Strip trailing \r\n from body
            if body.endswith(b"\r\n"):
                body = body[:-2]
            # Parse Content-Disposition to get field name
            cd = ""
            for hline in raw_headers.decode("utf-8", errors="replace").split("\r\n"):
                if hline.lower().startswith("content-disposition"):
                    cd = hline
                    break
            name = None
            for item in cd.split(";"):
                item = item.strip()
                if item.lower().startswith("name="):
                    name = item[5:].strip('"')
            if name:
                fields[name] = body

    confirm = (fields.get("confirm") or b"false").decode() == "true"
    new_bytes = fields.get("file")
    if not new_bytes:
        handler.json_response(400, {"error": "Geen bestand gevonden in de upload."})
        return
    dest_path = DATA_DIR / cfg["filename"]

    if confirm:
        # Save file
        dest_path.write_bytes(new_bytes)
        handler.json_response(200, {"ok": True, "message": f"{cfg['label']} opgeslagen."})
        return

    # Preview mode: compute diff
    old_bytes = dest_path.read_bytes() if dest_path.exists() else None
    try:
        headers, added, removed, unchanged = _diff_xlsx(
            old_bytes, new_bytes, cfg["sheet"], cfg["key_cols"]
        )
    except Exception as e:
        handler.json_response(500, {"error": f"Fout bij verwerken van bestand: {e}"})
        return

    handler.json_response(200, {
        "ok": True,
        "preview": True,
        "label": cfg["label"],
        "headers": headers,
        "added": added[:200],     # cap at 200 rows for the preview
        "removed": removed[:50],
        "unchanged": unchanged,
        "total_new": len(added),
        "total_removed": len(removed),
    })


def main():
    init_db()
    server = ThreadingHTTPServer((HOST, PORT), Handler)
    print(f"Planner running on http://{HOST}:{PORT}")
    server.serve_forever()


if __name__ == "__main__":
    main()
